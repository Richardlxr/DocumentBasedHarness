from __future__ import annotations

import os
import shutil
import tempfile
from collections.abc import Iterable
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.table import _Cell

from ..diagrams.artifacts import (
    DiagramArtifactBuilder,
    DiagramExportProfile,
    stable_diagram_stem,
)
from ..diagrams.model import DiagramConverterRegistry, DiagramSource
from ..errors import DocumentError, SourceLocation
from ..extension import ExtensionRegistry
from ..ir import (
    BlockQuote,
    CodeBlock,
    DiagramBlock,
    DocumentIR,
    Heading,
    HorizontalRule,
    InlineImage,
    InlineMath,
    LineBreak,
    ListBlock,
    MathBlock,
    Paragraph,
    TableBlock,
    TextSpan,
)
from ..lifecycle import RenderContext, RenderHooks
from ..math_conversion import tex_to_omml
from ..table_builder import TableCellSpec, TableRowSpec, TableSpec, build_table
from ..table_format import PLAIN_COMPACT_TABLE_PROFILE, TableFormatProfile
from ..template import configure_document
from .diagram_ooxml import add_png_picture, add_vsdx_object
from .math_ooxml import append_omml
from .ooxml import (
    add_hyperlink,
    add_paragraph_bottom_border,
    clear_cell,
    create_numbering_instance,
    set_paragraph_numbering,
)
from .styles import (
    BODY_STYLE,
    BODY_TEXT_STYLE,
    BULLET_STYLE,
    CODE_BLOCK_STYLE,
    CODE_CHARACTER_STYLE,
    HEADING_STYLES,
    MATH_BLOCK_STYLE,
    NUMBER_STYLE,
    QUOTE_STYLE,
    TABLE_BULLET_STYLE,
    TABLE_CODE_BLOCK_STYLE,
    TABLE_CODE_CHARACTER_STYLE,
    TABLE_HEADING_STYLE,
    TABLE_MATH_BLOCK_STYLE,
    TABLE_NUMBER_STYLE,
    TABLE_STYLE,
    TABLE_TEXT_STYLE,
)


class DocxRenderer:
    def __init__(
        self,
        *,
        registry: ExtensionRegistry,
        template: str | Path | None = None,
        source_dir: str | Path = ".",
        table_profile: TableFormatProfile | None = None,
        diagram_profile: DiagramExportProfile | None = None,
        diagram_registry: DiagramConverterRegistry | None = None,
        hooks: RenderHooks | None = None,
        source: str = "<string>",
        project_root: str | Path | None = None,
    ) -> None:
        self.registry = registry
        self.template = Path(template) if template else None
        self.source_dir = Path(source_dir)
        self.table_profile = table_profile or PLAIN_COMPACT_TABLE_PROFILE
        self.diagram_profile = diagram_profile or DiagramExportProfile()
        self.diagram_builder = DiagramArtifactBuilder(
            registry=diagram_registry,
            profile=self.diagram_profile,
            source_dir=self.source_dir,
        )
        self.diagram_output_dir = Path(".")
        self.diagram_index = 0
        self.document = None
        self.hooks = hooks or RenderHooks()
        self.source = source
        self.project_root = Path(project_root).resolve() if project_root is not None else None

    def render(self, document_ir: DocumentIR, output: str | Path) -> Path:
        if self.template:
            if not self.template.exists():
                raise DocumentError(f"template does not exist: {self.template}")
            document = Document(self.template)
        else:
            document = Document()
            configure_document(document)

        destination = Path(output)
        destination.parent.mkdir(parents=True, exist_ok=True)
        context = RenderContext(
            source=self.source,
            source_dir=self.source_dir,
            output=destination.resolve(),
            project_root=self.project_root,
        )
        with tempfile.TemporaryDirectory(
            prefix=f".{destination.stem}-build-", dir=destination.parent
        ) as temporary:
            staging = Path(temporary)
            staged_output = staging / destination.name
            staged_diagrams = staging / "diagrams" / destination.stem
            self.document = document
            self.diagram_output_dir = staged_diagrams
            self.hooks.configure(document, context)
            # Pass the public Document container. Its add_table() wrapper supplies
            # the writable page width; the internal _Body method requires callers
            # to calculate that implementation detail themselves.
            self.render_blocks(document_ir.blocks, document)
            self.hooks.finalize(document, context)
            self.hooks.validate(document, context)
            document.save(staged_output)
            self._publish_build(staged_output, staged_diagrams, destination)
        return destination

    @staticmethod
    def _publish_build(staged_output: Path, staged_diagrams: Path, destination: Path) -> None:
        final_diagrams = destination.parent / "diagrams" / destination.stem
        final_diagrams.parent.mkdir(parents=True, exist_ok=True)
        backup = final_diagrams.with_name(f".{final_diagrams.name}.backup-{uuid4().hex}")
        had_previous = final_diagrams.exists()
        installed_new = False
        try:
            if had_previous:
                os.replace(final_diagrams, backup)
            if staged_diagrams.exists():
                os.replace(staged_diagrams, final_diagrams)
                installed_new = True
            os.replace(staged_output, destination)
        except Exception:
            if installed_new and final_diagrams.exists():
                shutil.rmtree(final_diagrams)
            if had_previous and backup.exists():
                os.replace(backup, final_diagrams)
            raise
        finally:
            if backup.exists():
                shutil.rmtree(backup)

    def render_blocks(self, blocks: Iterable[object], container) -> None:
        for block in blocks:
            self.render_block(block, container)

    def render_block(self, block: object, container) -> None:
        if self.registry.render_block(block, self, container):
            return

        if isinstance(block, Heading):
            style = (
                TABLE_HEADING_STYLE if isinstance(container, _Cell) else HEADING_STYLES[block.level]
            )
            paragraph = self.add_paragraph(container, style)
            self.render_inlines(paragraph, block.inlines, block.location)
            return

        if isinstance(block, Paragraph):
            paragraph = self.add_paragraph(container)
            self.render_inlines(paragraph, block.inlines, block.location)
            return

        if isinstance(block, CodeBlock):
            style = TABLE_CODE_BLOCK_STYLE if isinstance(container, _Cell) else CODE_BLOCK_STYLE
            paragraph = self.add_paragraph(container, style)
            paragraph.add_run(block.text)
            return

        if isinstance(block, MathBlock):
            style = TABLE_MATH_BLOCK_STYLE if isinstance(container, _Cell) else MATH_BLOCK_STYLE
            paragraph = self.add_paragraph(container, style)
            append_omml(
                paragraph,
                tex_to_omml(block.tex, display=True, location=block.location),
            )
            return

        if isinstance(block, DiagramBlock):
            self.render_diagram(block, container)
            return

        if isinstance(block, BlockQuote):
            for child in block.blocks:
                if isinstance(child, Paragraph):
                    style = TABLE_TEXT_STYLE if isinstance(container, _Cell) else QUOTE_STYLE
                    paragraph = self.add_paragraph(container, style)
                    self.render_inlines(paragraph, child.inlines, child.location)
                else:
                    self.render_block(child, container)
            return

        if isinstance(block, ListBlock):
            self.render_list(block, container)
            return

        if isinstance(block, TableBlock):
            self.render_table(block, container)
            return

        if isinstance(block, HorizontalRule):
            paragraph = self.add_paragraph(container)
            add_paragraph_bottom_border(paragraph)
            return

        raise DocumentError(f"no DOCX renderer registered for IR node: {type(block).__name__}")

    def render_diagram(self, block: DiagramBlock, container) -> None:
        self.diagram_index += 1
        artifact = self.diagram_builder.build(
            DiagramSource(block.source_kind, block.content, block.location),
            self.diagram_output_dir,
            stem=stable_diagram_stem(block.content, self.diagram_index),
        )
        paragraph = self.add_paragraph(container, BODY_STYLE)
        # Fixed line spacing in formal-document body styles can clip inline OLE previews.
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        available_width = self.available_width_twips(container)
        available_height = self.available_height_twips()
        if self.diagram_profile.embed_format == "vsdx":
            add_vsdx_object(
                paragraph,
                artifact,
                available_width_twips=available_width,
                available_height_twips=available_height,
                profile=self.diagram_profile,
            )
        else:
            add_png_picture(
                paragraph,
                artifact,
                available_width_twips=available_width,
                available_height_twips=available_height,
                profile=self.diagram_profile,
            )

    def render_list(self, block: ListBlock, container) -> None:
        if isinstance(container, _Cell):
            style = TABLE_NUMBER_STYLE if block.ordered else TABLE_BULLET_STYLE
        else:
            style = NUMBER_STYLE if block.ordered else BULLET_STYLE
        num_id = create_numbering_instance(self.document, style) if block.ordered else None
        for item in block.items:
            first, *rest = item
            if isinstance(first, Paragraph):
                paragraph = self.add_paragraph(container, style)
                self.render_inlines(paragraph, first.inlines, first.location)
            else:
                paragraph = self.add_paragraph(container, style)
                self.render_block(first, container)
            if num_id is not None:
                set_paragraph_numbering(paragraph, num_id)
            self.render_blocks(rest, container)

    def render_table(self, block: TableBlock, container) -> None:
        rows = ([block.header] if block.header else []) + list(block.rows)
        if not rows:
            return
        column_count = max(cell.column + cell.column_span for row in rows for cell in row.cells)

        def cell_writer(cell_data, *, heading: bool):
            def write(cell) -> None:
                self.render_blocks(cell_data.blocks, cell)
                if heading:
                    for paragraph in cell.paragraphs:
                        paragraph.style = TABLE_HEADING_STYLE

            return write

        row_specs = tuple(
            TableRowSpec(
                cells=tuple(
                    TableCellSpec(
                        column=cell_data.column,
                        column_span=cell_data.column_span,
                        row_span=cell_data.row_span,
                        write=cell_writer(
                            cell_data,
                            heading=block.header is not None and row_index == 0,
                        ),
                    )
                    for cell_data in row.cells
                ),
                repeat_header=block.header is not None and row_index == 0,
            )
            for row_index, row in enumerate(rows)
        )
        build_table(
            container,
            TableSpec(columns=column_count, rows=row_specs, style=TABLE_STYLE),
            available_width_twips=self.available_width_twips(container),
            profile=self.table_profile,
        )
        self.add_spacing_after(container)

    def render_inlines(
        self,
        paragraph,
        inlines,
        location: SourceLocation | None = None,
    ) -> None:
        for inline in inlines:
            if isinstance(inline, TextSpan):
                if inline.href:
                    if inline.href.startswith("#"):
                        raise DocumentError(
                            "internal document links are not supported yet",
                            location,
                        )
                    add_hyperlink(
                        paragraph,
                        inline.text,
                        inline.href,
                        bold=inline.bold,
                        italic=inline.italic,
                    )
                    continue
                run = paragraph.add_run(inline.text)
                run.bold = inline.bold
                run.italic = inline.italic
                if inline.code:
                    run.style = (
                        TABLE_CODE_CHARACTER_STYLE
                        if isinstance(paragraph._parent, _Cell)
                        else CODE_CHARACTER_STYLE
                    )
            elif isinstance(inline, InlineImage):
                image_path = (self.source_dir / inline.uri).resolve()
                if not image_path.exists():
                    raise DocumentError(f"image does not exist: {inline.uri}", location)
                paragraph.add_run().add_picture(str(image_path))
            elif isinstance(inline, InlineMath):
                append_omml(
                    paragraph,
                    tex_to_omml(inline.tex, display=False, location=inline.location),
                )
            elif isinstance(inline, LineBreak):
                paragraph.add_run().add_break(WD_BREAK.LINE)
            else:
                raise DocumentError(
                    f"unsupported inline IR node: {type(inline).__name__}", location
                )

    def prepare_cell(self, cell) -> None:
        clear_cell(cell)

    def available_width_twips(self, container) -> int:
        if isinstance(container, _Cell) and container.width is not None:
            return max(1, int(container.width) // 635)
        section = self.document.sections[-1]
        width = section.page_width - section.left_margin - section.right_margin
        return max(1, int(width) // 635)

    def available_height_twips(self) -> int:
        section = self.document.sections[-1]
        height = section.page_height - section.top_margin - section.bottom_margin
        return max(1, int(height) // 635)

    def add_paragraph(self, container, style: str | None = None):
        if style is None:
            style = TABLE_TEXT_STYLE if isinstance(container, _Cell) else BODY_TEXT_STYLE
        if isinstance(container, _Cell):
            paragraphs = container.paragraphs
            has_content = any(child.tag != qn("w:pPr") for child in paragraphs[0]._p)
            if len(paragraphs) == 1 and not has_content:
                paragraph = paragraphs[0]
                paragraph.style = style
                return paragraph
        return container.add_paragraph(style=style)

    def add_spacing_after(self, container) -> None:
        if not isinstance(container, _Cell):
            container.add_paragraph(style=BODY_STYLE)
