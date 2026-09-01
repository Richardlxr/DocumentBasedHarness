from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .errors import DocumentError


def _length_cm(value) -> float | None:
    return round(value.cm, 3) if value is not None else None


def _style_font(style) -> dict[str, object | None]:
    font = style.font
    color = font.color.rgb
    east_asia = None
    r_pr = style._element.rPr
    if r_pr is not None and r_pr.rFonts is not None:
        east_asia = r_pr.rFonts.get(qn("w:eastAsia"))
    return {
        "name": font.name,
        "east_asia": east_asia,
        "size_pt": round(font.size.pt, 2) if font.size is not None else None,
        "bold": font.bold,
        "italic": font.italic,
        "color": str(color) if color is not None else None,
    }


def _style_record(style) -> dict[str, object | None]:
    return {
        "name": style.name,
        "style_id": style.style_id,
        "type": style.type.name.lower(),
        "builtin": getattr(style, "builtin", None),
        "base_style": style.base_style.name if style.base_style else None,
        "font": _style_font(style),
    }


def _table_record(table, index: int) -> dict[str, object]:
    cells = []
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            tc_pr = cell._tc.tcPr
            grid_span = None
            vertical_merge = None
            if tc_pr is not None:
                if tc_pr.gridSpan is not None:
                    grid_span = int(tc_pr.gridSpan.val)
                if tc_pr.vMerge is not None:
                    vertical_merge = str(tc_pr.vMerge.val or "continue")
            cells.append(
                {
                    "row": row_index,
                    "column": column_index,
                    "text": cell.text[:160],
                    "grid_span": grid_span,
                    "vertical_merge": vertical_merge,
                }
            )

    has_images = bool(table._tbl.xpath(".//a:blip"))
    has_hyperlinks = bool(table._tbl.xpath(".//w:hyperlink"))
    return {
        "index": index,
        "component": f"components/table-{index + 1:03d}.docx",
        "style": table.style.name if table.style is not None else None,
        "rows": len(table.rows),
        "columns": len(table.columns),
        "column_widths_emu": [column.width for column in table.columns],
        "has_images": has_images,
        "has_hyperlinks": has_hyperlinks,
        "cells": cells,
    }


def analyze_docx(source: str | Path) -> dict[str, object]:
    source_path = Path(source)
    if not source_path.exists():
        raise DocumentError(f"DOCX style source does not exist: {source_path}")
    document = Document(source_path)
    paragraph_styles = [
        _style_record(style)
        for style in document.styles
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER)
    ]
    table_styles = [
        _style_record(style) for style in document.styles if style.type == WD_STYLE_TYPE.TABLE
    ]
    tables = [_table_record(table, index) for index, table in enumerate(document.tables)]
    warnings = []
    for table in tables:
        if table["has_images"] or table["has_hyperlinks"]:
            warnings.append(
                f"{table['component']} contains relationships; use it as a visual reference "
                "or reimplement those relationships explicitly."
            )

    sections = []
    for section in document.sections:
        sections.append(
            {
                "page_width_cm": _length_cm(section.page_width),
                "page_height_cm": _length_cm(section.page_height),
                "top_margin_cm": _length_cm(section.top_margin),
                "bottom_margin_cm": _length_cm(section.bottom_margin),
                "left_margin_cm": _length_cm(section.left_margin),
                "right_margin_cm": _length_cm(section.right_margin),
            }
        )

    return {
        "schema": "docx-harness/style-manifest/v1",
        "source": source_path.name,
        "sha256": hashlib.sha256(source_path.read_bytes()).hexdigest(),
        "sections": sections,
        "paragraph_and_character_styles": paragraph_styles,
        "table_styles": table_styles,
        "tables": tables,
        "warnings": warnings,
    }


def write_style_manifest(source: str | Path, output: str | Path) -> Path:
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(analyze_docx(source), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return destination


def _clear_body(document) -> None:
    body = document._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def _write_table_components(source: Path, destination: Path) -> None:
    source_document = Document(source)
    table_xml = [deepcopy(table._tbl) for table in source_document.tables]
    destination.mkdir(parents=True, exist_ok=True)
    for index, table in enumerate(table_xml, start=1):
        component = Document(source)
        _clear_body(component)
        body = component._element.body
        section_properties = body.sectPr
        insertion_index = (
            body.index(section_properties) if section_properties is not None else len(body)
        )
        body.insert(insertion_index, table)
        body.insert(insertion_index + 1, OxmlElement("w:p"))
        component.save(destination / f"table-{index:03d}.docx")


def extract_template(
    source: str | Path,
    output: str | Path,
    *,
    manifest: str | Path | None = None,
    components: str | Path | None = None,
) -> Path:
    source_path = Path(source).resolve()
    destination = Path(output)
    if source_path == destination.resolve():
        raise DocumentError("style source and extracted template must be different files")
    document = Document(source_path)
    _clear_body(document)
    document.core_properties.creator = "docx-harness"
    document.core_properties.subject = "Template extracted from a user DOCX"
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)

    if manifest is not None:
        write_style_manifest(source_path, manifest)
    if components is not None:
        _write_table_components(source_path, Path(components))
    return destination
