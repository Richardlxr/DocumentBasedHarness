from __future__ import annotations

from dataclasses import dataclass, replace
from typing import Literal

from docutils import nodes
from docutils.parsers.rst import Directive, directives
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from ..errors import SourceLocation
from ..extension import DirectiveExtension
from ..ir import Heading, Inline, ListBlock, Paragraph, TextSpan
from ..renderers.styles import (
    TABLE_FIELD_LABEL_STYLE,
    TEST_CASE_TABLE_STYLE,
    TEST_CASE_TITLE_STYLE,
)
from ..table_builder import TableCellSpec, TableRowSpec, TableSpec, build_table

_TEST_CASE_COLUMN_WEIGHTS = (0.65, 0.3, 1.8, 0.65, 0.3, 1.8)


class TestCaseMeta(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    id: str = Field(pattern=r"^[A-Z][A-Z0-9_-]*$")
    priority: Literal["P0", "P1", "P2", "P3"] = "P2"
    schema_version: Literal["testcase/v1"] = Field(default="testcase/v1", alias="schema")
    title: str = Field(min_length=1)


class TestCaseNode(nodes.General, nodes.Element):
    pass


class TestCaseDirective(Directive):
    required_arguments = 1
    final_argument_whitespace = True
    has_content = True
    option_spec = {
        "id": directives.unchanged_required,
        "priority": directives.unchanged,
        "schema": directives.unchanged,
    }

    def run(self):
        values = {
            "id": self.options.get("id"),
            "priority": self.options.get("priority", "P2"),
            "schema": self.options.get("schema", "testcase/v1"),
            "title": self.arguments[0].strip(),
        }
        try:
            meta = TestCaseMeta.model_validate(values)
        except ValidationError as error:
            raise self.error(str(error)) from error

        node = TestCaseNode()
        node["meta"] = meta.model_dump(by_alias=True)
        node.line = self.lineno
        self.state.nested_parse(self.content, self.content_offset, node)
        return [node]


@dataclass(frozen=True, slots=True)
class TestCaseBlock:
    meta: TestCaseMeta
    blocks: tuple[object, ...]
    location: SourceLocation
    preamble: tuple[object, ...] = ()
    sections: tuple[TestCaseSection, ...] = ()


@dataclass(frozen=True, slots=True)
class TestCaseSection:
    title: tuple[Inline, ...]
    blocks: tuple[object, ...]
    location: SourceLocation | None = None


def _trim_leading_space(inlines: tuple[Inline, ...]) -> tuple[Inline, ...]:
    remaining = list(inlines)
    while remaining and isinstance(remaining[0], TextSpan):
        first = remaining[0]
        trimmed = first.text.lstrip()
        if trimmed:
            remaining[0] = replace(first, text=trimmed)
            break
        remaining.pop(0)
    while remaining and isinstance(remaining[-1], TextSpan) and not remaining[-1].text:
        remaining.pop()
    return tuple(remaining)


def _section_boundary(
    block: object,
) -> tuple[tuple[Inline, ...], Paragraph | None, SourceLocation | None] | None:
    if isinstance(block, Heading):
        return block.inlines, None, block.location
    if not isinstance(block, Paragraph) or not block.inlines:
        return None
    first_index = next(
        (
            index
            for index, inline in enumerate(block.inlines)
            if not isinstance(inline, TextSpan) or inline.text
        ),
        None,
    )
    if first_index is None:
        return None
    first = block.inlines[first_index]
    if not isinstance(first, TextSpan) or not first.bold:
        return None
    label = first.text.strip()
    if not label.endswith((":", "：")):
        return None
    label = label[:-1].strip()
    if not label:
        return None
    title = (replace(first, text=label, bold=False),)
    remainder = _trim_leading_space(block.inlines[first_index + 1 :])
    paragraph = Paragraph(remainder, block.location) if remainder else None
    return title, paragraph, block.location


def _partition_blocks(
    blocks: tuple[object, ...],
) -> tuple[tuple[object, ...], tuple[TestCaseSection, ...]]:
    preamble: list[object] = []
    sections: list[TestCaseSection] = []
    title: tuple[Inline, ...] | None = None
    content: list[object] = []
    location: SourceLocation | None = None

    def finish_section() -> None:
        if title is not None:
            sections.append(TestCaseSection(title, tuple(content), location))

    for block in blocks:
        boundary = _section_boundary(block)
        if boundary is None:
            (preamble if title is None else content).append(block)
            continue
        finish_section()
        title, leading_paragraph, location = boundary
        content = [leading_paragraph] if leading_paragraph is not None else []
    finish_section()
    return tuple(preamble), tuple(sections)


def transform_test_case(node, transformer, heading_level: int) -> TestCaseBlock:
    blocks = tuple(transformer.transform_children(node, heading_level))
    preamble, sections = _partition_blocks(blocks)
    return TestCaseBlock(
        meta=TestCaseMeta.model_validate(node["meta"]),
        blocks=blocks,
        preamble=preamble,
        sections=sections,
        location=transformer.location(node),
    )


def render_test_case(block: TestCaseBlock, renderer, container) -> None:
    def text_writer(
        text: str,
        style: str | None = None,
        *,
        centered: bool = False,
    ):
        def write(cell) -> None:
            paragraph = renderer.add_paragraph(cell, style)
            paragraph.add_run(text)
            if centered:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        return write

    def inline_writer(inlines: tuple[Inline, ...], style: str | None = None):
        def write(cell) -> None:
            paragraph = renderer.add_paragraph(cell, style)
            renderer.render_inlines(paragraph, inlines)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        return write

    def content_writer(blocks: tuple[object, ...]):
        def write(cell) -> None:
            renderer.render_blocks(blocks, cell)
            if not blocks:
                renderer.add_paragraph(cell).add_run("（无内容）")

        return write

    def fallback_writer(cell) -> None:
        renderer.render_blocks(block.blocks, cell)
        if not block.blocks:
            renderer.add_paragraph(cell).add_run("（无内容）")

    rows: list[TableRowSpec] = [
        TableRowSpec(
            cells=(
                TableCellSpec(
                    column=0,
                    column_span=6,
                    write=text_writer(
                        block.meta.title,
                        TEST_CASE_TITLE_STYLE,
                        centered=True,
                    ),
                ),
            ),
            repeat_header=True,
        ),
        TableRowSpec(
            cells=(
                TableCellSpec(
                    column=0,
                    write=text_writer("用例 ID", TABLE_FIELD_LABEL_STYLE),
                ),
                TableCellSpec(column=1, column_span=2, write=text_writer(block.meta.id)),
                TableCellSpec(
                    column=3,
                    write=text_writer("优先级", TABLE_FIELD_LABEL_STYLE),
                ),
                TableCellSpec(column=4, column_span=2, write=text_writer(block.meta.priority)),
            ),
            repeat_header=True,
        ),
    ]

    if block.sections:
        if block.preamble:
            rows.append(
                TableRowSpec(
                    cells=(
                        TableCellSpec(
                            column=0,
                            column_span=6,
                            write=content_writer(block.preamble),
                        ),
                    ),
                    prevent_split=False,
                )
            )
        for section in block.sections:
            if len(section.blocks) == 1 and isinstance(section.blocks[0], ListBlock):
                list_block = section.blocks[0]
                for item_index, item in enumerate(list_block.items):
                    cells = [
                        TableCellSpec(
                            column=1,
                            write=text_writer(
                                f"{item_index + 1}." if list_block.ordered else "•",
                                centered=True,
                            ),
                        ),
                        TableCellSpec(
                            column=2,
                            column_span=4,
                            write=content_writer(item),
                        ),
                    ]
                    if item_index == 0:
                        cells.insert(
                            0,
                            TableCellSpec(
                                column=0,
                                row_span=len(list_block.items),
                                write=inline_writer(section.title, TABLE_FIELD_LABEL_STYLE),
                            ),
                        )
                    rows.append(TableRowSpec(cells=tuple(cells)))
            else:
                rows.append(
                    TableRowSpec(
                        cells=(
                            TableCellSpec(
                                column=0,
                                write=inline_writer(section.title, TABLE_FIELD_LABEL_STYLE),
                            ),
                            TableCellSpec(
                                column=1,
                                column_span=5,
                                write=content_writer(section.blocks),
                            ),
                        )
                    )
                )
    else:
        rows.append(
            TableRowSpec(
                cells=(TableCellSpec(column=0, column_span=6, write=fallback_writer),),
                prevent_split=False,
            )
        )

    profile = renderer.table_profile
    if profile.layout is not None:
        profile = replace(
            profile,
            layout=replace(
                profile.layout,
                column_weights=_TEST_CASE_COLUMN_WEIGHTS,
            ),
        )

    spec = TableSpec(
        columns=6,
        style=TEST_CASE_TABLE_STYLE,
        rows=tuple(rows),
    )
    build_table(
        container,
        spec,
        available_width_twips=renderer.available_width_twips(container),
        profile=profile,
    )
    renderer.add_spacing_after(container)


TEST_CASE_EXTENSION = DirectiveExtension(
    name="test-case",
    directive=TestCaseDirective,
    node_type=TestCaseNode,
    ir_type=TestCaseBlock,
    transform=transform_test_case,
    render=render_test_case,
)
