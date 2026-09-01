from __future__ import annotations

from dataclasses import dataclass

from docutils import nodes
from docutils.parsers.rst import Directive, directives
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.shared import Pt
from docx.table import _Cell

from ..errors import DocumentError, SourceLocation
from ..extension import DirectiveExtension
from ..renderers.styles import BODY_STYLE, HEADING_STYLES


class DocumentCoverNode(nodes.General, nodes.Element):
    pass


class DocumentCoverDirective(Directive):
    has_content = False
    option_spec = {"date": directives.unchanged_required}

    def run(self):
        date = self.options.get("date")
        if not date:
            raise self.error(":date: is required")
        node = DocumentCoverNode()
        node["date"] = date.strip()
        node.line = self.lineno
        return [node]


@dataclass(frozen=True, slots=True)
class DocumentCoverBlock:
    date: str
    location: SourceLocation


def transform_document_cover(node, transformer, heading_level):
    del heading_level
    return DocumentCoverBlock(
        date=node["date"],
        location=transformer.location(node),
    )


def _line_height_pt(style, fallback: float) -> float:
    spacing = style.paragraph_format.line_spacing
    if hasattr(spacing, "pt"):
        return float(spacing.pt)
    size = style.font.size
    font_size = float(size.pt) if size is not None else fallback
    if isinstance(spacing, (int, float)):
        return max(font_size, font_size * float(spacing))
    return max(font_size, font_size * 1.2)


def render_document_cover(block, renderer, container) -> None:
    if isinstance(container, _Cell):
        raise DocumentError("document covers are not supported inside table cells", block.location)
    if len(container.paragraphs) != 1 or container.tables:
        raise DocumentError(
            "document-cover must immediately follow the document H1", block.location
        )

    title = container.paragraphs[0]
    if title.style.name != HEADING_STYLES[1]:
        raise DocumentError(
            "document-cover must immediately follow the document H1", block.location
        )

    section = renderer.document.sections[-1]
    title_style = renderer.document.styles[HEADING_STYLES[1]]
    title_line_height = _line_height_pt(title_style, 24)
    page_height = float(section.page_height.pt)
    top_margin = float(section.top_margin.pt)
    title_space_before = max(0.0, page_height / 2 - top_margin - title_line_height / 2)

    title_format = title.paragraph_format
    title_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_format.first_line_indent = Pt(0)
    title_format.space_before = Pt(title_space_before)
    title_format.space_after = Pt(0)
    title_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    title_format.line_spacing = Pt(title_line_height)
    title_format.keep_with_next = False

    section.different_first_page_header_footer = True
    footer = section.first_page_footer
    footer.is_linked_to_previous = False
    date = footer.paragraphs[0]
    date.clear()
    date.style = BODY_STYLE
    date_format = date.paragraph_format
    date_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_format.first_line_indent = Pt(0)
    date_format.space_before = Pt(0)
    date_format.space_after = Pt(0)
    date.add_run(block.date)

    page_break = container.add_paragraph(style=BODY_STYLE)
    break_format = page_break.paragraph_format
    break_format.first_line_indent = Pt(0)
    break_format.space_before = Pt(0)
    break_format.space_after = Pt(0)
    break_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    break_format.line_spacing = Pt(1)
    page_break.add_run().add_break(WD_BREAK.PAGE)


DOCUMENT_COVER_EXTENSION = DirectiveExtension(
    name="document-cover",
    directive=DocumentCoverDirective,
    node_type=DocumentCoverNode,
    ir_type=DocumentCoverBlock,
    transform=transform_document_cover,
    render=render_document_cover,
)
