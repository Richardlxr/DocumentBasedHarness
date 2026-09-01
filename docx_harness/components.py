from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table

from .errors import DocumentError


def clone_table_component(container, component: str | Path, *, table_index: int = 0) -> Table:
    """Clone a relationship-free table prototype into a document or cell."""
    component_path = Path(component)
    document = Document(component_path)
    if table_index < 0 or table_index >= len(document.tables):
        raise DocumentError(f"table component {component_path} has no table at index {table_index}")
    source_table = document.tables[table_index]
    if source_table._tbl.xpath(".//a:blip") or source_table._tbl.xpath(".//w:hyperlink"):
        raise DocumentError(
            f"table component {component_path} contains image or hyperlink relationships; "
            "render those relationships explicitly"
        )

    cloned_xml = deepcopy(source_table._tbl)
    placeholder = container.add_table(rows=1, cols=1)
    placeholder._tbl.addnext(cloned_xml)
    placeholder._element.getparent().remove(placeholder._element)
    return Table(cloned_xml, container)
