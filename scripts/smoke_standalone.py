"""Exercise extracted executables through real gates with explicitly synthetic replies.

This fixture is not a user run or reader acceptance. It checks packaging at the
render boundary, including importlib metadata, diagram fonts and math resources.
"""

from __future__ import annotations

import json
import subprocess
import tempfile
from pathlib import Path

import yaml
from docx import Document
from pptx import Presentation


def smoke_render(bundle: Path) -> None:
    suffix = ".exe" if (bundle / "comh.exe").exists() else ""
    executable = bundle / f"comh{suffix}"
    with tempfile.TemporaryDirectory(prefix="release-render-smoke-") as temporary:
        root = Path(temporary) / "fixture"

        def cli(*args):
            completed = subprocess.run(
                [str(executable), *args],
                cwd=bundle,
                check=False,
                capture_output=True,
                text=True,
                encoding="utf-8",
                timeout=180,
            )
            if completed.returncode:
                raise RuntimeError(
                    f"release fixture command {args} failed:\n"
                    f"{completed.stdout}\n{completed.stderr}"
                )
            return completed.stdout

        def command(*args):
            return cli(*args, "--run", str(root))

        def write(path, data):
            file = root / path
            file.parent.mkdir(parents=True, exist_ok=True)
            file.write_text(
                yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8"
            )

        def accept(target):
            request = json.loads(command("present", target, "--json"))
            command(
                "respond",
                request["id"],
                "--decision",
                "accepted",
                "--reply",
                "Synthetic release fixture accepts this test version",
                "--source",
                "test:release-fixture",
            )

        cli("init-run", str(root))
        write(
            ".workspace/intake.yaml",
            {
                "goal": "Synthetic release rendering test",
                "source_scope": "Fixture data",
                "basis": "test:release-fixture",
            },
        )
        command("intake", str(root / ".workspace/intake.yaml"))
        (root / "sources/data.txt").write_text("Synthetic baseline: 12 ms\n", encoding="utf-8")
        write(
            "evidence/evidence.yaml",
            {
                "version": 1,
                "sources": [{"id": "SRC01", "path": "sources/data.txt"}],
                "items": [
                    {
                        "id": "E001",
                        "kind": "datum",
                        "content": "Synthetic baseline: 12 ms",
                        "source": {"source": "SRC01", "locator": "line=1"},
                        "value": {"number": 12, "unit": "ms"},
                    }
                ],
            },
        )
        write(
            "evidence/coverage.yaml",
            {
                "sources": [
                    {
                        "file": p.relative_to(root).as_posix(),
                        "status": "read",
                        "locator": "entire synthetic fixture",
                    }
                    for p in sorted((root / "sources").rglob("*"))
                    if p.is_file()
                ]
            },
        )
        command("save", "evidence")
        brief = {
            "language": "en",
            "audience": {"description": "Release test fixture"},
            "objective": "Check extracted runtime rendering",
            "takeaways": ["Baseline value"],
            "media": [
                {"medium": "pptx", "surface": "presentation"},
                {"medium": "html", "surface": "presentation"},
                {"medium": "docx", "surface": "report"},
            ],
            "delivery_context": "Synthetic offline test",
            "constraints": {"hard": [], "soft": []},
            "voice": {"style": "plain"},
            "visual_materials": "Native objects only",
            "appearance": {"selection": "default", "review": "sample"},
            "presentation": {"setting": "general"},
        }
        brief["alignment"] = {
            key: {
                "source": "user",
                "status": "provided",
                "basis": "test:release-fixture",
                "value": value,
            }
            for key, value in brief.items()
        }
        brief["version"] = 1
        write("brief/brief.yaml", brief)
        command("save", "brief")
        accept("brief")
        write(
            "narrative/narrative.yaml",
            {
                "version": 1,
                "claims": [
                    {
                        "id": "C01",
                        "statement": "Baseline is 12 ms",
                        "evidence": ["E001"],
                        "status": "supported",
                    }
                ],
                "story": [
                    {
                        "id": "S01",
                        "purpose": "demonstrate_effect",
                        "message": "Baseline is 12 ms",
                        "claims": ["C01"],
                    }
                ],
            },
        )
        command("save", "narrative")
        accept("narrative")
        write(
            "projection/deck_plan.yaml",
            {
                "version": 1,
                "deck": {
                    "title": "Synthetic release fixture",
                    "pages": [
                        {
                            "id": "P01",
                            "page_role": "content",
                            "beat": "S01",
                            "title": "Baseline value",
                            "visual": {
                                "table": {
                                    "columns": ["Case", "Time"],
                                    "rows": [["Baseline", {"value_from": "E001"}]],
                                }
                            },
                        },
                        {
                            "id": "P02",
                            "page_role": "content",
                            "beat": "S01",
                            "title": "Measurement path",
                            "visual": {
                                "diagram": {"mermaid": "flowchart LR\n a[Input] --> b[Measure]"}
                            },
                        },
                    ],
                },
            },
        )
        command("save", "deck_plan")
        accept("deck_outline")
        command("render", "deck")
        command("render", "deck-html")
        write(
            "projection/report_plan.yaml",
            {
                "version": 1,
                "sections": [
                    {
                        "id": "R01",
                        "heading": "Baseline",
                        "beats": ["S01"],
                        "claims": ["C01"],
                        "evidence": ["E001"],
                        "must_include": ["Synthetic measurement"],
                    }
                ],
            },
        )
        command("save", "report_plan")
        accept("report_outline")
        (root / "documents/report.md").write_text(
            "# Release fixture\n\n## Baseline\n\n"
            "Synthetic measurement: baseline is 12 ms. [E001]\n\n"
            "Symbolic expression: $x+y$.\n",
            encoding="utf-8",
        )
        command("save", "report_md")
        command("render", "report")
        command("validate", "all")
        prs = Presentation(root / "build/deck.pptx")
        assert len(prs.slides) == 2
        assert next(s.table for s in prs.slides[0].shapes if s.has_table).cell(1, 1).text == "12 ms"
        assert any(s.name.startswith("diagram-node:") for s in prs.slides[1].shapes)
        assert "<svg" in (root / "build/deck.html").read_text(encoding="utf-8")
        doc = Document(root / "build/report.docx")
        assert any("Synthetic measurement" in p.text for p in doc.paragraphs)
        assert "oMath" in doc.element.xml
        print(
            "Extracted bundle rendered editable PPTX, SVG HTML and DOCX math through fixture gates"
        )
