"""Object inventory, not a claim that raster images contain no useful evidence."""

from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..artifacts import Finding


def audit_editability(prs, page_ids):
    pages, findings = [], []
    area = prs.slide_width * prs.slide_height
    for page_id, slide in zip(page_ids, prs.slides, strict=True):
        counts = dict(text=0, tables=0, charts=0, connectors=0, pictures=0)
        largest = 0

        def visit(shapes, counts=counts):
            nonlocal largest
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(shape.shapes)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    counts["pictures"] += 1
                    largest = max(largest, shape.width * shape.height / area)
                elif shape.has_table:
                    counts["tables"] += 1
                elif shape.has_chart:
                    counts["charts"] += 1
                elif shape.shape_type == MSO_SHAPE_TYPE.LINE:
                    counts["connectors"] += 1
                elif shape.has_text_frame and shape.text.strip():
                    counts["text"] += 1

        visit(slide.shapes)
        pages.append(dict(id=page_id, **counts, largest_picture_fraction=round(largest, 3)))
        if largest >= 0.5:
            findings.append(
                Finding(
                    "deck_plan",
                    "editability:raster-dominant",
                    "warn",
                    "fail",
                    f"page {page_id}: one picture covers {largest:.0%} of the slide; "
                    "inspect whether authored text/table/diagram was flattened. "
                    "Source photos/screenshots may be valid.",
                    "deck_plan",
                )
            )
    return {
        "pages": pages,
        "scope": "native object counts; inspect actual editing and rendering",
    }, findings
