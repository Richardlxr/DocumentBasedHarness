"""Evidence tables: identical cell values in editable PPTX and semantic HTML.

Cells are plain text, ``{text, evidence}``, ``{value_from}`` (numbers pulled
from evidence, right-aligned), ``{status, tone, evidence}`` (an authored state
label such as 已验证) or ``{missing}`` (an authored reason a value is absent —
missing values stay visible). ``header_column`` marks row labels, ``highlight``
marks the rows/columns/cells the page is about, ``note`` states units,
conditions or source under the table. All emphasis is semantic; the theme
decides the colors.
"""

from __future__ import annotations

import html
from dataclasses import dataclass

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .metrics import fit_box, wrap_lines
from .ooxml import set_cell_borders, set_east_asian_font
from .text_forms import TONES, ink_on, status_style

_NOTE_GAP = 0.08


@dataclass(frozen=True, slots=True)
class Cell:
    text: str
    kind: str = "text"  # text | number | status | missing
    tone: str = "neutral"


def table_errors(spec: dict) -> list[str]:
    """Structure the renderer consumes; evidence references are checked by validate."""
    errors = []
    columns = spec.get("columns", [])
    rows = spec.get("rows", [])
    if any(len(row) != len(columns) for row in rows):
        errors.append("every table row must match its column count")
    weights = spec.get("column_weights")
    if weights is not None and (len(weights) != len(columns) or any(v <= 0 for v in weights)):
        errors.append("table column_weights must be positive and match columns")
    for row in rows:
        for cell in row:
            if not isinstance(cell, dict):
                continue
            if "status" in cell:
                if not str(cell.get("status") or "").strip():
                    errors.append("table status cells need a visible status label")
                if cell.get("tone", "neutral") not in TONES:
                    errors.append(
                        f"table status tone '{cell.get('tone')}' is not one of {', '.join(TONES)}"
                    )
            if "missing" in cell and not str(cell.get("missing") or "").strip():
                errors.append("table missing cells need visible wording (e.g. 未测)")
    highlight = spec.get("highlight") or {}
    for index in highlight.get("rows", []):
        if not (isinstance(index, int) and 0 <= index < len(rows)):
            errors.append(f"table highlight row {index} is out of range ({len(rows)} rows)")
    for index in highlight.get("columns", []):
        if not (isinstance(index, int) and 0 <= index < len(columns)):
            errors.append(
                f"table highlight column {index} is out of range ({len(columns)} columns)"
            )
    for pair in highlight.get("cells", []):
        if not (
            isinstance(pair, list)
            and len(pair) == 2
            and all(isinstance(v, int) for v in pair)
            and 0 <= pair[0] < len(rows)
            and 0 <= pair[1] < len(columns)
        ):
            errors.append(f"table highlight cell {pair} is out of range")
    return errors


def table_cells(spec, evidence, page_id) -> list[list[Cell]]:
    """Header row first; body cells resolved against evidence."""
    from .deck import _evidence_number, _fmt

    errors = table_errors(spec)
    if errors:
        raise ValueError(f"page {page_id}: {'; '.join(errors)}")
    items = {item["id"]: item for item in (evidence or {}).get("items", [])}
    result = [[Cell(str(label)) for label in spec["columns"]]]
    for row in spec["rows"]:
        values = []
        for cell in row:
            if isinstance(cell, dict) and "value_from" in cell:
                number, unit = _evidence_number(items, cell["value_from"], page_id, "table cell")
                values.append(Cell(_fmt(number) + (f" {unit}" if unit else ""), "number"))
            elif isinstance(cell, dict) and "status" in cell:
                values.append(Cell(str(cell["status"]), "status", cell.get("tone", "neutral")))
            elif isinstance(cell, dict) and "missing" in cell:
                values.append(Cell(str(cell["missing"]), "missing"))
            else:
                values.append(Cell(str(cell["text"] if isinstance(cell, dict) else cell)))
        result.append(values)
    return result


def table_rows(spec, evidence, page_id) -> list[list[str]]:
    return [[cell.text for cell in row] for row in table_cells(spec, evidence, page_id)]


def _highlighted(spec: dict, body_row: int, column: int) -> bool:
    marks = spec.get("highlight") or {}
    return (
        body_row in marks.get("rows", [])
        or column in marks.get("columns", [])
        or [body_row, column] in marks.get("cells", [])
    )


def _numeric_columns(cells: list[list[Cell]]) -> set[int]:
    body = cells[1:]
    return {
        j
        for j in range(len(cells[0]))
        if body
        and any(row[j].kind == "number" for row in body)
        and all(row[j].kind in {"number", "missing"} for row in body)
    }


@dataclass(frozen=True, slots=True)
class TableLayout:
    widths: list[float]
    heights: list[float]
    note_h: float

    @property
    def table_h(self) -> float:
        return sum(self.heights)


def layout_table(spec, evidence, page_id, theme, *, w: float, h: float) -> TableLayout:
    """Measured column widths and row heights at the fixed body size, or a
    source-oriented ValueError. The validator runs this as a dry run."""
    cells = table_cells(spec, evidence, page_id)
    weights = spec.get("column_weights") or [1] * len(cells[0])
    widths = [w * v / sum(weights) for v in weights]
    size = theme.theme.body_size
    heights = [
        max(
            len(wrap_lines(cell.text, size, width * 72 - 14, theme.body_font))
            for cell, width in zip(row, widths, strict=False)
        )
        * size
        * 1.25
        / 72
        + 0.16
        for row in cells
    ]
    note = str(spec.get("note") or "")
    note_h = 0.0
    if note:
        caption = theme.theme.caption_size
        note_h = len(wrap_lines(note, caption, w * 72 - 8, theme.body_font)) * caption * 1.3 / 72
        note_h += 0.1 + _NOTE_GAP
    if sum(heights) + note_h > h:
        raise ValueError(
            f"page {page_id}: table does not fit at {size}pt; split rows or simplify text"
        )
    for i, row in enumerate(cells):
        for j, cell in enumerate(row):
            fit = fit_box(
                cell.text,
                width_pt=widths[j] * 72 - 14,
                height_pt=heights[i] * 72 - 10,
                max_size=size,
                min_size=size,
                family=theme.body_font,
            )
            if fit.overflows:
                raise ValueError(f"page {page_id}: table cell [{i},{j}] does not fit")
    return TableLayout(widths, heights, note_h)


def render_table(slide, spec, evidence, page_id, theme, *, x, y, w, h) -> list[int]:
    """Draw the table (and its note); returns the shape ids of the visual."""
    t = theme.theme
    cells = table_cells(spec, evidence, page_id)
    layout = layout_table(spec, evidence, page_id, theme, w=w, h=h)
    numeric = _numeric_columns(cells)
    size = t.body_size
    frame = slide.shapes.add_table(
        len(cells), len(cells[0]), Inches(x), Inches(y), Inches(w), Inches(layout.table_h)
    )
    frame.name = "evidence-table"
    table = frame.table
    table.first_row = True
    table.first_col = bool(spec.get("header_column"))
    table.horz_banding = False
    for column, width in zip(table.columns, layout.widths, strict=False):
        column.width = Inches(width)
    for i, (row, height) in enumerate(zip(cells, layout.heights, strict=False)):
        table.rows[i].height = Inches(height)
        for j, value in enumerate(row):
            header = i == 0
            row_label = not header and j == 0 and bool(spec.get("header_column"))
            marked = not header and _highlighted(spec, i - 1, j)
            fill, color = None, "text"
            if value.kind == "status":
                fill = status_style(value.tone)["fill"]
            elif marked:
                fill = "accent_soft"
            elif row_label:
                fill = "card_fill"
            if fill and fill != "card_fill":
                color = ink_on(theme, fill)
            if value.kind == "missing":
                color = "muted"
            cell = table.cell(i, j)
            cell.margin_left = cell.margin_right = Pt(7)
            cell.margin_top = cell.margin_bottom = Pt(5)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = getattr(t, fill)
            else:
                cell.fill.background()
            cell.text = value.text
            cell.text_frame.word_wrap = True
            last = i == len(cells) - 1
            set_cell_borders(
                cell,
                bottom=None if last else ((t.accent, 2.0) if header else (t.card_line, 0.75)),
            )
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.15
                p.space_after = p.space_before = Pt(0)
                if j in numeric:
                    p.alignment = PP_ALIGN.RIGHT
                elif value.kind == "status":
                    p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.bold = header or row_label or marked or value.kind == "status"
                    run.font.italic = value.kind == "missing"
                    run.font.name = theme.body_font
                    run.font.color.rgb = getattr(t, color)
                    set_east_asian_font(run, theme.body_font)
    ids = [frame.shape_id]
    note = str(spec.get("note") or "")
    if note:
        box = slide.shapes.add_textbox(
            Inches(x),
            Inches(y + layout.table_h + _NOTE_GAP),
            Inches(w),
            Inches(layout.note_h - _NOTE_GAP),
        )
        box.name = "table-note"
        box.text_frame.word_wrap = True
        box.text_frame.margin_left = box.text_frame.margin_right = Pt(2)
        box.text_frame.margin_top = box.text_frame.margin_bottom = Pt(2)
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = note
        for run in paragraph.runs:
            run.font.size = Pt(t.caption_size)
            run.font.name = theme.body_font
            run.font.color.rgb = t.muted
            set_east_asian_font(run, theme.body_font)
        ids.append(box.shape_id)
    return ids


def table_html(spec, evidence, page_id):
    cells = table_cells(spec, evidence, page_id)
    numeric = _numeric_columns(cells)
    weights = spec.get("column_weights") or [1] * len(cells[0])
    columns = "".join(f'<col style="width:{v / sum(weights) * 100:.4f}%">' for v in weights)
    output = [f'<table class="evidence-table"><colgroup>{columns}</colgroup>']
    for i, row in enumerate(cells):
        rendered = []
        for j, cell in enumerate(row):
            classes = []
            if j in numeric:
                classes.append("num")
            if i and _highlighted(spec, i - 1, j):
                classes.append("hl")
            text = html.escape(cell.text)
            if cell.kind == "status":
                text = f'<span class="form-pill tone-{cell.tone}">{text}</span>'
            elif cell.kind == "missing":
                text = f'<em class="muted">{text}</em>'
            attr = f' class="{" ".join(classes)}"' if classes else ""
            if i == 0:
                rendered.append(f'<th scope="col"{attr}>{text}</th>')
            elif j == 0 and spec.get("header_column"):
                rendered.append(f'<th scope="row"{attr}>{text}</th>')
            else:
                rendered.append(f"<td{attr}>{text}</td>")
        output.append("<tr>" + "".join(rendered) + "</tr>")
    output.append("</table>")
    note = str(spec.get("note") or "")
    if note:
        output.append(f'<p class="table-note muted">{html.escape(note)}</p>')
    return "".join(output)
