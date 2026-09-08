"""Editable Mermaid flowcharts using the compiler's public, measured layout API.

This adapter consumes generated mxGraph XML, not arbitrary draw.io documents.
PPTX and HTML share the same nodes, routed segments and label positions.
"""

from __future__ import annotations

import html
import math
from dataclasses import dataclass

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from docx_harness import DiagramSource, DiagramStyleProfile, create_mermaid_flowchart_converter

from .metrics import fit_box, text_width
from .ooxml import set_connector_arrow, set_east_asian_font


@dataclass
class DiagramScene:
    nodes: list[dict]
    edges: list[dict]
    bounds: tuple[float, float, float, float]


def _style(cell) -> dict:
    value = cell.get("style", "")
    style = dict(v.split("=", 1) for v in value.split(";") if "=" in v)
    if "rhombus" in value.split(";"):
        style["shape"] = "rhombus"
    return style


def _along(points, fraction):
    lengths = [math.dist(a, b) for a, b in zip(points, points[1:], strict=False)]
    remaining = sum(lengths) * fraction
    for a, b, length in zip(points, points[1:], lengths, strict=False):
        if length and remaining <= length:
            return tuple(a[i] + (b[i] - a[i]) * remaining / length for i in (0, 1))
        remaining -= length
    return points[-1]


def diagram_scene(source: str, family: str = "Microsoft YaHei") -> DiagramScene:
    document = create_mermaid_flowchart_converter(DiagramStyleProfile(font_family=family))(
        DiagramSource("mermaid", source)
    )
    nodes, edges, boxes = [], [], []
    for cell in document.graph_model.findall("./root/mxCell"):
        if cell.get("vertex") != "1":
            continue
        g = cell.find("mxGeometry")
        box = tuple(float(g.get(k, 0)) for k in ("x", "y", "width", "height"))
        node = dict(id=cell.get("id"), text=cell.get("value", ""), box=box, style=_style(cell))
        nodes.append(node)
        boxes.append(box)
    by_id = {n["id"]: n for n in nodes}
    for cell in document.graph_model.findall("./root/mxCell"):
        if cell.get("edge") != "1":
            continue
        style = _style(cell)
        endpoints = []
        for key, prefix in (("source", "exit"), ("target", "entry")):
            x, y, w, h = by_id[cell.get(key)]["box"]
            endpoints.append(
                (x + w * float(style[prefix + "X"]), y + h * float(style[prefix + "Y"]))
            )
        points = (
            [endpoints[0]]
            + [
                (float(p.get("x")), float(p.get("y")))
                for p in cell.findall("./mxGeometry/Array/mxPoint")
            ]
            + [endpoints[1]]
        )
        edge = dict(id=cell.get("id"), points=points, style=style, text=cell.get("value", ""))
        if edge["text"]:
            g = cell.find("mxGeometry")
            cx, cy = _along(points, (float(g.get("x", 0)) + 1) / 2)
            offset = g.find("mxPoint[@as='offset']")
            if offset is not None:
                cx += float(offset.get("x", 0))
                cy += float(offset.get("y", 0))
            size = int(style["fontSize"])
            lines = edge["text"].splitlines()
            w = max(text_width(t, size, family) for t in lines) + 8
            h = len(lines) * size * 1.25 + 4
            edge["box"] = (cx - w / 2, cy - h / 2, w, h)
            boxes.append(edge["box"])
        boxes.extend((x, y, 0, 0) for x, y in points)
        edges.append(edge)
    left = min(b[0] for b in boxes) - 4
    top = min(b[1] for b in boxes) - 4
    right = max(b[0] + b[2] for b in boxes) + 4
    bottom = max(b[1] + b[3] for b in boxes) + 4
    return DiagramScene(nodes, edges, (left, top, right - left, bottom - top))


def render_native_diagram(slide, scene, theme, *, x, y, w, h):
    """Return native objects; never silently fall back to a bitmap or shrink to illegibility."""
    bx, by, bw, bh = scene.bounds
    scale = min(w / bw, h / bh)
    ox, oy = x + (w - bw * scale) / 2, y + (h - bh * scale) / 2
    created = []

    def text_shape(obj, shape):
        style, value = obj["style"], obj["text"]
        size = float(style["fontSize"]) * scale * 72
        if size < 17:
            raise ValueError(
                "native diagram text falls below 17pt; use a full-width visual or split it"
            )
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Pt(2)
        frame.margin_top = frame.margin_bottom = Pt(1)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        width, height = shape.width / 12700 - 4, shape.height / 12700 - 2
        fit = fit_box(
            value,
            width_pt=width,
            height_pt=height,
            min_size=math.ceil(size),
            max_size=math.ceil(size),
            family=theme.body_font,
        )
        if fit.overflows:
            raise ValueError(f"native diagram {obj['id']}: text does not fit: {value!r}")
        from pptx.enum.text import PP_ALIGN

        p = frame.paragraphs[0]
        p.text = value
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.15
        p.space_after = p.space_before = Pt(0)
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.name = theme.body_font
            run.font.bold = bool(int(style.get("fontStyle", 0)) & 1)
            run.font.color.rgb = RGBColor.from_string(style["fontColor"].lstrip("#"))
            set_east_asian_font(run, theme.body_font)

    def position(box):
        a, b, c, d = box
        return [
            Inches(v) for v in (ox + (a - bx) * scale, oy + (b - by) * scale, c * scale, d * scale)
        ]

    for edge in scene.edges:
        for i, (a, b) in enumerate(zip(edge["points"], edge["points"][1:], strict=False)):
            if a == b:
                continue
            shape = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(ox + (a[0] - bx) * scale),
                Inches(oy + (a[1] - by) * scale),
                Inches(ox + (b[0] - bx) * scale),
                Inches(oy + (b[1] - by) * scale),
            )
            shape.name = f"diagram-edge:{edge['id']}:{i}"
            style = edge["style"]
            shape.line.color.rgb = RGBColor.from_string(style["strokeColor"].lstrip("#"))
            shape.line.width = Pt(float(style["strokeWidth"]) * scale * 72)
            if style.get("dashed") == "1":
                shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if i == len(edge["points"]) - 2 and style.get("endArrow") != "none":
                set_connector_arrow(shape)
            if i == 0 and style.get("startArrow", "none") != "none":
                set_connector_arrow(shape, at_start=True)
            created.append(shape)
    for node in scene.nodes:
        style = node["style"]
        kind = {
            "rhombus": MSO_SHAPE.DIAMOND,
            "ellipse": MSO_SHAPE.OVAL,
            "cylinder3": MSO_SHAPE.CAN,
        }.get(
            style.get("shape"),
            MSO_SHAPE.ROUNDED_RECTANGLE if style.get("rounded") == "1" else MSO_SHAPE.RECTANGLE,
        )
        shape = slide.shapes.add_shape(kind, *position(node["box"]))
        shape.name = f"diagram-node:{node['id']}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(style["fillColor"].lstrip("#"))
        shape.line.color.rgb = RGBColor.from_string(style["strokeColor"].lstrip("#"))
        if style.get("dashed") == "1":
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        text_shape(node, shape)
        created.append(shape)
    for edge in scene.edges:
        if not edge["text"]:
            continue
        shape = slide.shapes.add_textbox(*position(edge["box"]))
        shape.name = f"diagram-label:{edge['id']}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(
            edge["style"].get("labelBackgroundColor", "#FFFFFF").lstrip("#")
        )
        text_shape(edge, shape)
        created.append(shape)
    return created


def diagram_svg(scene: DiagramScene) -> str:
    """Inline SVG keeps HTML labels searchable and shares the native routing."""
    bx, by, bw, bh = scene.bounds
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" role="img" '
        f'viewBox="{bx} {by} {bw} {bh}" style="width:100%;max-height:440px" '
        f'font-family="{html.escape(scene.nodes[0]["style"]["fontFamily"], quote=True)}">'
    ]
    for edge in scene.edges:
        style = edge["style"]
        points = " ".join(f"{x},{y}" for x, y in edge["points"])
        dash = ' stroke-dasharray="5 4"' if style.get("dashed") == "1" else ""
        parts.append(
            f'<polyline points="{points}" fill="none" stroke="{style["strokeColor"]}" '
            f'stroke-width="{style["strokeWidth"]}"{dash}/>'
        )
        arrow_ends = []
        if style.get("startArrow", "none") != "none":
            arrow_ends.append((edge["points"][1], edge["points"][0]))
        if style.get("endArrow", "none") != "none":
            arrow_ends.append(tuple(edge["points"][-2:]))
        for a, b in arrow_ends:
            angle = math.atan2(b[1] - a[1], b[0] - a[0])
            pts = [b] + [
                (b[0] - 7 * math.cos(angle + d), b[1] - 7 * math.sin(angle + d))
                for d in (-0.45, 0.45)
            ]
            parts.append(
                '<polygon points="'
                + " ".join(f"{x},{y}" for x, y in pts)
                + f'" fill="{style["strokeColor"]}"/>'
            )
    for obj in scene.nodes + [e for e in scene.edges if e["text"]]:
        x, y, w, h = obj["box"]
        st = obj["style"]
        fill = st.get("fillColor", st.get("labelBackgroundColor", "#FFFFFF"))
        stroke = st.get("strokeColor", "none") if obj in scene.nodes else "none"
        if st.get("shape") == "rhombus":
            points = f"{x + w / 2},{y} {x + w},{y + h / 2} {x + w / 2},{y + h} {x},{y + h / 2}"
            parts.append(f'<polygon points="{points}" fill="{fill}" stroke="{stroke}"/>')
        elif st.get("shape") == "ellipse":
            parts.append(
                f'<ellipse cx="{x + w / 2}" cy="{y + h / 2}" rx="{w / 2}" ry="{h / 2}" '
                f'fill="{fill}" stroke="{stroke}"/>'
            )
        elif st.get("shape") == "cylinder3":
            parts.append(
                f'<path d="M{x},{y + 6} C{x},{y - 2} {x + w},{y - 2} {x + w},{y + 6} '
                f'L{x + w},{y + h - 6} C{x + w},{y + h + 2} {x},{y + h + 2} {x},{y + h - 6} Z" '
                f'fill="{fill}" stroke="{stroke}"/>'
            )
            parts.append(
                f'<ellipse cx="{x + w / 2}" cy="{y + 6}" rx="{w / 2}" ry="6" '
                f'fill="{fill}" stroke="{stroke}"/>'
            )
        else:
            parts.append(
                f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" stroke="{stroke}"/>'
            )
        lines, size = obj["text"].splitlines(), float(st["fontSize"])
        for i, line in enumerate(lines):
            ty = y + h / 2 + (i - (len(lines) - 1) / 2) * size * 1.15
            parts.append(
                f'<text x="{x + w / 2}" y="{ty}" text-anchor="middle" '
                f'dominant-baseline="middle" font-size="{size}" '
                f'font-weight="{"bold" if int(st.get("fontStyle", 0)) & 1 else "normal"}" '
                f'fill="{st["fontColor"]}">{html.escape(line)}</text>'
            )
    return "".join(parts) + "</svg>"
