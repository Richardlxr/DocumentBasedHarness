"""Resolve appearance to typed IR, run the existing content layout, compose, then publish."""

from __future__ import annotations

import math
import os
from dataclasses import dataclass, replace
from pathlib import Path

from ..render.metrics import fixed_text_sizes, font_resolution
from ..render.theme import (
    _SIZE_FIELDS,
    background_is_light,
    contrast_ratio,
    fonts_for,
    merge_tokens,
    resolve_style,
)
from .compose import _box, compose
from .package import NS, Package, StyleError, descendants, select_shapes, shape_id, shapes
from .profile import StyleTemplate, Surface, load_style
from .text import assess_text


@dataclass(frozen=True, slots=True)
class PageStyleIR:
    page_id: str
    role: str
    surface: Surface


@dataclass(frozen=True, slots=True)
class StyleDeckIR:
    template: StyleTemplate
    pages: tuple[PageStyleIR, ...]
    width: int
    height: int


def _require_light_background(source: Package, surface: Surface, width: int, height: int):
    """Never treat a theme's lt1 token as proof about native artwork or a photograph."""
    from pptx.dml.color import RGBColor

    scopes = source.scopes(surface.slide)
    native = None
    artwork_area = 0.0
    for scope in ("slide", "layout", "master"):
        root = source.tree(scopes[scope])
        if native is None:
            native = root.find("p:cSld/p:bg", NS)
        for shape in select_shapes(shapes(root), surface.selected(scope)):
            box = _box(shape)
            if box:
                artwork_area += (box[2] - box[0]) * (box[3] - box[1])
            if artwork_area >= width * height * 0.5:
                raise StyleError(
                    "light-background constraint cannot be certified for large native artwork; "
                    "use an inspected plain-light surface"
                )
    if native is None:
        return  # PresentationML default background is white.
    color = native.find("p:bgPr/a:solidFill/a:srgbClr", NS)
    if color is None or len(color):
        raise StyleError(
            "light-background constraint cannot be certified for an image, theme fill "
            "or transformed native background"
        )
    if not background_is_light(RGBColor.from_string(color.get("val"))):
        raise StyleError("imported appearance conflicts with light-background constraint")


def resolve_plan(plan: dict, run_root: Path, *, allow_dark: bool = True) -> StyleDeckIR:
    style = plan.get("deck", {}).get("style", {})
    template = load_style(run_root, style.get("pptx_style"))
    if style.get("template"):
        raise StyleError("choose pptx_style or the legacy token template, not both")
    source = Package(template.source)
    size = source.tree("ppt/presentation.xml").find("p:sldSz", NS)
    width, height = (int(size.get(k)) for k in ("cx", "cy"))
    if width <= 0 or height <= 0:
        raise StyleError("style template has invalid canvas dimensions")
    # Existing layouts are validated for these canvases. Reject unsupported sizes explicitly.
    if not (abs(width / height - 16 / 9) < 0.01 or abs(width / height - 4 / 3) < 0.01):
        raise StyleError("style compiler currently supports 16:9 and 4:3 canvases")
    pages = []
    base = resolve_style(style, run_root=run_root).theme
    for page in plan.get("deck", {}).get("pages", []):
        role = page.get("page_role", "content")
        surface = template.surface(role)
        if page.get("background") or (page.get("visual") or {}).get("background"):
            raise StyleError(
                f"{page['id']}: page background conflicts with the appearance template"
            )
        theme = merge_tokens(merge_tokens(base, surface.tokens), style.get("tokens_override"))
        if not allow_dark:
            _require_light_background(source, surface, width, height)
        if contrast_ratio(theme.text, theme.card_fill) < 4.5:
            raise StyleError(f"{page['id']}: body/card contrast is below 4.5:1")
        pages.append(PageStyleIR(page["id"], role, surface))
    if not pages:
        raise StyleError("style compilation needs at least one page")
    return StyleDeckIR(template, tuple(pages), width, height)


def compile_deck(
    plan: dict,
    run_root: Path,
    output: Path,
    *,
    language: str,
    evidence: dict | None,
    allow_dark: bool,
):
    from ..render.deck import render_deck
    from ..render.theme import RenderTheme

    run_root = run_root.resolve()
    ir = resolve_plan(plan, run_root, allow_dark=allow_dark)
    if output.resolve().is_relative_to(ir.template.path.parent):
        raise StyleError("compiled output must not overwrite template inputs")
    base = resolve_style(plan["deck"].get("style"), run_root=run_root).theme
    canvas_w, canvas_h = 12191695, 6858000  # existing 13.333 × 7.5 inch layout space
    scale = min(ir.width / canvas_w, ir.height / canvas_h)
    dy = round((ir.height - canvas_h * scale) / 2)
    themes, font_maps = [], []
    for page in ir.pages:
        theme = merge_tokens(
            merge_tokens(base, page.surface.tokens), plan["deck"]["style"].get("tokens_override")
        )
        changes, font_map = {}, {}
        for attribute in _SIZE_FIELDS.values():
            size = getattr(theme, attribute)
            virtual = math.ceil(size / scale)
            changes[attribute] = virtual
            font_map[virtual * 100] = size * 100
        theme = replace(theme, **changes)
        title_font, body_font = fonts_for(theme, language)
        themes.append(RenderTheme(theme, title_font, body_font))
        font_maps.append(font_map)
    stage = run_root / ".workspace" / "pptx-style-compile"
    if stage.resolve() != stage or any(
        (stage / n).is_symlink() for n in ("active.lock", "content.pptx", "assembled.pptx")
    ):
        raise StyleError("style staging path must not traverse symlinks")
    stage.mkdir(parents=True, exist_ok=True)
    lock = stage / "active.lock"
    try:
        handle = lock.open("x", encoding="utf-8")
    except FileExistsError as error:
        raise StyleError(
            f"another style build is active, or interrupted: inspect {lock}"
        ) from error
    native, assembled = stage / "content.pptx", stage / "assembled.pptx"
    try:
        handle.write(str(os.getpid()))
        handle.close()
        with fixed_text_sizes():
            result = render_deck(
                plan,
                run_root,
                native,
                language=language,
                evidence=evidence,
                allow_dark=allow_dark,
                _appearance_themes=themes,
            )
        failures = [f.detail for f in result.findings if f.verdict == "fail"]
        if failures:
            raise StyleError("style content layout failed: " + "; ".join(failures))
        source, dest = Package(ir.template.source), Package(native)
        native_text = []
        for page in ir.pages:
            for scope, part in source.scopes(page.surface.slide).items():
                for node in descendants(
                    select_shapes(shapes(source.tree(part)), page.surface.selected(scope))
                ):
                    assessment = assess_text(source, page.surface.slide, scope, node)
                    if assessment is not None:
                        native_text.append(
                            {
                                "page": page.page_id,
                                "scope": scope,
                                "shape": shape_id(node),
                                "single_line": shape_id(node)
                                in dict(page.surface.single_line).get(scope, ()),
                                **assessment,
                            }
                        )
        mapping = compose(
            source,
            dest,
            [p.surface for p in ir.pages],
            [p.page_id for p in ir.pages],
            font_maps,
            scale,
            dy,
        )
        dest.save(assembled)
        # Validate packaging with the public reader before publishing the new file.
        from pptx import Presentation

        parsed = Presentation(assembled)
        if len(parsed.slides) != len(ir.pages):
            raise StyleError("compiled slide count does not match the content plan")
        for page, slide in zip(ir.pages, parsed.slides, strict=True):
            ids = [shape.shape_id for shape in slide.shapes]
            if len(ids) != len(set(ids)):
                raise StyleError(f"{page.page_id}: duplicate output shape IDs")
            for shape in slide.shapes:
                if shape.name.startswith(("diagram-node:", "diagram-label:")):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.size is not None and run.font.size.pt < 17:
                                raise StyleError(
                                    f"{page.page_id}: native diagram text falls below 17pt "
                                    "on the template canvas; simplify or split the diagram"
                                )
        result.metadata = {
            "compiler": "comh/pptx-style/v1",
            "style": ir.template.name,
            "source_sha256": ir.template.source_hash,
            "profile_sha256": ir.template.profile_hash,
            "canvas_emu": [ir.width, ir.height],
            "pages": [
                {
                    "id": p.page_id,
                    "appearance_surface": p.surface.role,
                    "source_slide": p.surface.slide,
                    "kept": {scope: list(ids) for scope, ids in p.surface.keep},
                    "single_line": {scope: list(ids) for scope, ids in p.surface.single_line},
                }
                for p in ir.pages
            ],
            "native_assets": mapping["copied_assets"],
            "native_text_diagnostics": native_text,
            "typography": [
                {"title": font_resolution(t.title_font), "body": font_resolution(t.body_font)}
                for t in themes
            ],
            "layout_policy": (
                "existing content layout; source slots/density not imported; "
                "native style supplies footer furniture"
            ),
            "visual_verification": "not performed here; inspect actual PPTX renders",
        }
        from ..render.editability import audit_editability

        result.metadata["editability"], edit_findings = audit_editability(
            parsed, [p.page_id for p in ir.pages]
        )
        result.findings = [f for f in result.findings if not f.check.startswith("editability:")]
        result.findings.extend(edit_findings)
        result.shape_map = {
            pid: {
                address: [
                    mapping["shape_ids"][pid][v] for v in values if v in mapping["shape_ids"][pid]
                ]
                for address, values in addresses.items()
            }
            for pid, addresses in result.shape_map.items()
        }
        output.parent.mkdir(parents=True, exist_ok=True)
        os.replace(assembled, output)
        result.output, result.theme = output, ir.template.name
        return result
    finally:
        handle.close()
        for path in (native, assembled, lock):
            path.unlink(missing_ok=True)
        stage.rmdir()
