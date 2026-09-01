"""Deck renderer tests: themes, measured layout, evidence-backed charts,
transitions, figure captions, language fonts."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation

from comh.render.deck import render_deck
from comh.render.theme import select_theme

PLAN = {
    "version": 1,
    "deck": {
        "title": "测试 Deck",
        "pages": [
            {
                "id": "P01",
                "page_role": "cover",
                "title": "测试 Deck",
                "support_points": ["副标题"],
                "notes": "开场",
            },
            {
                "id": "P02",
                "page_role": "content",
                "beat": "S01",
                "title": "标题应表达本页信息 18.2%",
                "support_points": ["点一", "点二"],
                "notes": "备注内容",
                "demotions": [{"content": "细节", "to": "notes"}],
            },
        ],
    },
}

EVIDENCE = {
    "items": [
        {"id": "E001", "kind": "datum", "content": "基线 220ms",
         "value": {"number": 220, "unit": "ms"}, "source": {"source": "SRC01", "locator": "r1"}},
        {"id": "E002", "kind": "datum", "content": "优化 180ms",
         "value": {"number": 180, "unit": "ms"}, "source": {"source": "SRC01", "locator": "r2"}},
    ]
}


def _texts(prs) -> str:
    return "\n".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


def test_render_deck_slides_notes_and_fonts(tmp_path: Path):
    result = render_deck(PLAN, tmp_path, tmp_path / "out.pptx", language="zh-CN")
    prs = Presentation(str(result.output))
    assert len(prs.slides) == 2
    assert prs.slides[1].notes_slide.notes_text_frame.text == "备注内容"
    assert "18.2%" in _texts(prs)
    assert "点一" in _texts(prs)
    fonts = {
        run.font.name
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    }
    assert fonts == {"Microsoft YaHei"}, "zh language drives CJK fonts"
    assert result.theme == "tier1-light"


def test_theme_selection_forces_light_when_dark_forbidden():
    choice = select_theme({"template": "midnight"}, allow_dark=False)
    assert choice.forced_light
    assert choice.theme.is_light
    assert select_theme({"template": "midnight"}, allow_dark=True).theme.name == "midnight"
    unknown = select_theme({"template": "no-such-theme"}, allow_dark=True)
    assert unknown.fallback_reason and unknown.theme.name == "tier1-light"


def test_chart_renders_from_evidence(tmp_path: Path):
    plan = {
        "version": 1,
        "deck": {
            "title": "t",
            "style": {"template": "slate-tech", "transition": "fade"},
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "带图页",
                    "support_points": ["点"],
                    "visual": {
                        "chart": {
                            "type": "column",
                            "title": "P99 (ms)",
                            "series": [
                                {"label": "baseline", "value_from": "E001"},
                                {"label": "partitioned", "value_from": "E002"},
                            ],
                        }
                    },
                }
            ],
        },
    }
    result = render_deck(plan, tmp_path, tmp_path / "chart.pptx", language="en", evidence=EVIDENCE)
    assert result.theme == "slate-tech"
    assert result.transition == "fade"
    assert not result.findings, [f.detail for f in result.findings]
    with ZipFile(result.output) as archive:
        names = archive.namelist()
        charts = [n for n in names if n.startswith("ppt/charts/chart")]
        slides = [n for n in names if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        assert charts, "chart part must exist"
        assert "<p:transition" in archive.read(slides[0]).decode("utf-8")
    prs = Presentation(str(result.output))
    chart_shapes = [sh for sh in prs.slides[0].shapes if sh.has_chart]
    assert chart_shapes
    plot = chart_shapes[0].chart.plots[0]
    assert [v for v in plot.series[0].values] == [220.0, 180.0], "values come from evidence"


def test_chart_missing_evidence_fails_loud(tmp_path: Path):
    import pytest

    plan = {
        "version": 1,
        "deck": {
            "title": "t",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "x",
                    "visual": {"chart": {"series": [{"value_from": "E999"}]}},
                }
            ],
        },
    }
    with pytest.raises(RuntimeError, match="E999"):
        render_deck(plan, tmp_path, tmp_path / "bad.pptx", evidence=EVIDENCE)


def test_overflow_produces_layout_finding(tmp_path: Path):
    long_title = "这是一个非常长的标题" * 20
    plan = {
        "version": 1,
        "deck": {
            "title": "t",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": long_title,
                    "support_points": ["点"],
                }
            ],
        },
    }
    result = render_deck(plan, tmp_path, tmp_path / "over.pptx", language="zh-CN")
    overflow = [f for f in result.findings if f.check == "layout"]
    assert overflow and "P01" in overflow[0].detail
    assert overflow[0].owning_artifact == "deck_plan"


def test_figure_with_caption_renders(tmp_path: Path):
    image = tmp_path / "assets" / "fig.png"
    image.parent.mkdir()
    Image.new("RGB", (60, 40), "white").save(image)
    plan = {
        "version": 1,
        "deck": {
            "title": "t",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "带图页",
                    "support_points": ["点"],
                    "visual": {
                        "asset_refs": [
                            {"ref": "assets/fig.png", "caption": "图注（E001）", "evidence": "E001"}
                        ]
                    },
                }
            ],
        },
    }
    result = render_deck(plan, tmp_path, tmp_path / "fig.pptx", language="zh-CN")
    prs = Presentation(str(result.output))
    slide = prs.slides[0]
    assert any(shape.shape_type == 13 for shape in slide.shapes), "picture embedded"
    assert any("图注（E001）" in sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    assert not result.findings


def test_elaborated_points_cards_and_callout(tmp_path: Path):
    plan = {
        "version": 1,
        "deck": {
            "title": "t",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "标题",
                    "support_points": [
                        {"point": "加粗导语", "detail": "展开说明一行"},
                        "普通要点",
                    ],
                    "metric_cards": [
                        {"label": "P99 降幅", "value_from": "E003"},
                        {"label": "基线", "value_from": "E001"},
                    ],
                    "callout": {"text": "底部结论条"},
                }
            ],
        },
    }
    evidence = {
        "items": EVIDENCE["items"]
        + [{"id": "E003", "kind": "datum", "content": "降幅 18.2%",
            "value": {"number": 18.2, "unit": "%"}, "source": {"source": "S", "locator": "x"}}]
    }
    result = render_deck(
        plan, tmp_path, tmp_path / "rich.pptx", language="zh-CN", evidence=evidence
    )
    assert not result.findings, [f.detail for f in result.findings]
    texts = [
        sh.text_frame.text
        for sh in Presentation(str(result.output)).slides[0].shapes
        if sh.has_text_frame
    ]
    joined = "\n".join(texts)
    assert "加粗导语" in joined and "展开说明一行" in joined and "普通要点" in joined
    assert "18.2 %" in joined and "220 ms" in joined, "card values pulled from evidence"
    assert "底部结论条" in joined
    # rounded rectangles: 2 metric cards + 2 point cards + 1 callout band
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

    def _rounded(shape) -> bool:
        try:
            return shape.auto_shape_type == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        except (ValueError, AttributeError):
            return False

    shapes = Presentation(str(result.output)).slides[0].shapes
    rounded = [sh for sh in shapes if _rounded(sh)]
    assert len(rounded) == 5


def test_metric_card_missing_evidence_fails_loud(tmp_path: Path):
    import pytest

    plan = {
        "version": 1,
        "deck": {
            "title": "t",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "x",
                    "metric_cards": [{"label": "l", "value_from": "E999"}],
                }
            ],
        },
    }
    with pytest.raises(RuntimeError, match="E999"):
        render_deck(plan, tmp_path, tmp_path / "bad.pptx", evidence=EVIDENCE)


def test_visual_validation_covers_cards_and_callout(tmp_path: Path):
    from comh.validate import validate_visuals

    artifacts = {
        "evidence": {"items": EVIDENCE["items"]},
        "deck_plan": {
            "deck": {
                "title": "t",
                "pages": [
                    {
                        "id": "P01",
                        "page_role": "content",
                        "title": "x",
                        "metric_cards": [{"label": "l", "value_from": "E001"},
                                          {"label": "bad", "value_from": "E999"}],
                        "callout": {"text": "c", "evidence": "E888"},
                    }
                ],
            }
        },
    }
    findings = validate_visuals(artifacts)
    details = " | ".join(f.detail for f in findings)
    assert "metric card references missing evidence 'E999'" in details
    assert "callout references missing evidence 'E888'" in details
    assert not any("'E001'" in f.detail for f in findings)


def test_cli_init_and_guard_flow(tmp_path: Path):
    from comh.cli import main

    assert main(["init-run", str(tmp_path / "run")]) == 0
    run = tmp_path / "run"
    assert (run / "run.yaml").is_file()
    assert (run / "templates" / "base.docx").is_file()
    assert main(["status", "--run", str(run)]) == 0
    # saving narrative before the brief gate must be refused (exit code 2)
    (run / "narrative" / "narrative.yaml").write_text(
        "version: 1\nclaims: []\nstory:\n  - {id: S01, purpose: p, message: m}\n",
        encoding="utf-8",
    )
    assert main(["save", "narrative", "--run", str(run)]) == 2
