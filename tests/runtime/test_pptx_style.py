"""Appearance compilation against native, editable Office fixtures; no content slots/density."""

from __future__ import annotations

import hashlib
from copy import deepcopy
from zipfile import ZipFile

import pytest
import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from comh.cli import main
from comh.pptx_style import StyleError, import_style, inspect_pptx, load_style
from comh.pptx_style.package import Package
from comh.render.deck import render_deck
from comh.render.html_deck import render_html_deck
from comh.scaffold import init_run


def fixture(tmp_path, *, ratio="16:9"):
    source = tmp_path / "brand.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333 if ratio == "16:9" else 10)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image = tmp_path / "logo.png"
    Image.new("RGB", (40, 40), "red").save(image)
    logo = slide.shapes.add_picture(str(image), Inches(0.1), Inches(0.1), Inches(0.2), Inches(0.2))
    logo.name = "ORIGINAL_LOGO"
    sample = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    sample.text = "SAMPLE CONTENT 98765"
    sample.text_frame.paragraphs[0].runs[0].font.size = Pt(37)
    slide.notes_slide.notes_text_frame.text = "SAMPLE PRIVATE NOTES"
    prs.slide_master.background.fill.solid()
    prs.slide_master.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    prs.save(source)
    run = tmp_path / "run"
    run.mkdir()
    path = import_style(source, "brand", run)
    data = yaml.safe_load(path.read_text())
    data["tokens"]["fonts"] = {"latin": ["Arial", "Arial"], "cjk": ["微软雅黑", "微软雅黑"]}
    data["tokens"]["sizes"] = {"cover_title": 40, "content_title": 32, "body_wide": 24}
    data["surfaces"] = {
        "content": {
            "slide": 1,
            "keep": {"master": [], "layout": [], "slide": [logo.shape_id]},
            "protect": {"slide": [logo.shape_id]},
        },
        "cover": {"slide": 1, "keep": {"master": [], "layout": [], "slide": [logo.shape_id]}},
    }
    path.write_text(yaml.safe_dump(data, allow_unicode=True))
    return source, run, path


def plan():
    return {
        "version": 1,
        "deck": {
            "title": "编译验证",
            "style": {"pptx_style": "templates/brand/style.yaml"},
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "样式与内容分离",
                    "support_points": ["保持原有内容结构"],
                    "notes": "USER NOTES",
                },
                {
                    "id": "P02",
                    "page_role": "content",
                    "title": "修改这一页",
                    "support_points": ["其他页面保持独立"],
                },
            ],
        },
    }


def test_single_line_wordmark_preserves_text_font_and_geometry(tmp_path):
    from comh.pptx_style.package import NS, blob

    _, run, path = fixture(tmp_path)
    source = path.parent / "source.pptx"
    prs = Presentation(source)
    label = prs.slides[0].shapes.add_textbox(Inches(0.5), Inches(6.5), Pt(61.23), Pt(19.63))
    label.text = "S J T U"
    label.text_frame.word_wrap = True
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    identity = label.shape_id
    prs.save(source)
    original = source.read_bytes()
    data = yaml.safe_load(path.read_text())
    data["sha256"] = hashlib.sha256(original).hexdigest()
    data["surfaces"]["content"]["keep"]["slide"].append(identity)
    data["surfaces"]["content"]["single_line"] = {"slide": [identity]}
    path.write_text(yaml.safe_dump(data))
    result = render_deck(plan(), run, run / "fixed.pptx")
    rebuilt = Presentation(result.output)
    fixed = next(s for s in rebuilt.slides[0].shapes if s.has_text_frame and s.text == "S J T U")
    assert fixed.text_frame.word_wrap is False
    assert len(fixed.text_frame.paragraphs) == 1
    assert fixed.text_frame.paragraphs[0].runs[0].font.size == Pt(18)
    assert (fixed.left, fixed.top, fixed.height) == (
        label.left,
        label.top,
        label.height,
    )
    assert fixed.width >= label.width
    assert result.metadata["native_text_diagnostics"][0]["single_line"]
    assert blob(fixed._element.find("p:txBody/a:p", NS)) == blob(
        label._element.find("p:txBody/a:p", NS)
    )
    assert source.read_bytes() == original
    # Existing authorial line breaks must never be silently removed by this repair.
    label.text = "S J T\nU"
    prs.save(source)
    data["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    path.write_text(yaml.safe_dump(data))
    with pytest.raises(StyleError, match="single_line requires one text paragraph"):
        render_deck(plan(), run, run / "multiline.pptx")


def test_single_line_cannot_target_omitted_content(tmp_path):
    _, run, path = fixture(tmp_path)
    data = yaml.safe_load(path.read_text())
    data["surfaces"]["content"]["single_line"] = {"slide": [999]}
    path.write_text(yaml.safe_dump(data))
    with pytest.raises(StyleError, match="single_line shapes must also be kept"):
        load_style(run, str(path.relative_to(run)))


@pytest.mark.parametrize("obstacle", ["canvas", "content", "decoration"])
def test_single_line_expansion_must_not_create_collisions(tmp_path, obstacle):
    _, run, path = fixture(tmp_path)
    source = path.parent / "source.pptx"
    prs = Presentation(source)
    page = prs.slides[0]
    x, y = (
        (13.2, 6.5) if obstacle == "canvas" else (0.5, 0.8) if obstacle == "content" else (0.5, 6.5)
    )
    label = page.shapes.add_textbox(Inches(x), Inches(y), Pt(3), Pt(20))
    label.text = "A C M E"
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    data = yaml.safe_load(path.read_text())
    keep = data["surfaces"]["content"]["keep"]["slide"]
    keep.append(label.shape_id)
    if obstacle == "decoration":
        other = page.shapes.add_picture(
            str(tmp_path / "logo.png"), Inches(0.8), Inches(y), Pt(15), Pt(20)
        )
        keep.append(other.shape_id)
    prs.save(source)
    data["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    data["surfaces"]["content"]["single_line"] = {"slide": [label.shape_id]}
    path.write_text(yaml.safe_dump(data))
    with pytest.raises(StyleError, match="leaves canvas|overlaps"):
        render_deck(plan(), run, run / "bad.pptx")


def test_native_text_risks_are_inventory_and_compile_diagnostics_without_auto_rewriting(tmp_path):
    _, run, path = fixture(tmp_path)
    source = path.parent / "source.pptx"
    prs = Presentation(source)
    label = prs.slides[0].shapes[1]
    label.width = Pt(5)
    label.text = "OTHER BRAND"
    label.text_frame.paragraphs[0].runs[0].font.name = "Missing Brand Font Fixture"
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    prs.save(source)
    data = yaml.safe_load(path.read_text())
    data["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    data["surfaces"]["content"]["keep"]["slide"].append(label.shape_id)
    path.write_text(yaml.safe_dump(data))
    inventory = inspect_pptx(source, slide=1)
    risk = inventory["slides"][0]["layers"]["slide"]["shapes"][1]["text_layout"]
    assert "width_pressure" in risk["risks"]
    assert "measurement_font_substitution" in risk["risks"]
    result = render_deck(plan(), run, run / "diagnosed.pptx")
    assert "width_pressure" in result.metadata["native_text_diagnostics"][0]["risks"]
    output = Presentation(result.output)
    copied = next(s for s in output.slides[0].shapes if s.has_text_frame and s.text == label.text)
    assert copied.width == label.width and copied.text_frame.word_wrap == label.text_frame.word_wrap


@pytest.mark.parametrize("ratio", ["16:9", "4:3"])
def test_native_style_preserves_assets_canvas_typography_and_content(tmp_path, ratio):
    source, run, path = fixture(tmp_path, ratio=ratio)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    original = Package(source)
    p = plan()
    result = render_deck(p, run, run / "out.pptx", language="zh-CN")
    prs = Presentation(result.output)
    assert len(prs.slides) == 2
    assert (prs.slide_width, prs.slide_height) == (Presentation(source).slide_width, Inches(7.5))
    assert result.metadata["compiler"] == "comh/pptx-style/v1"
    texts = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert "SAMPLE" not in texts and "98765" not in texts
    assert "样式与内容分离" in texts and "保持原有内容结构" in texts
    assert prs.slides[0].notes_slide.notes_text_frame.text == "USER NOTES"
    assert "SAMPLE PRIVATE" not in prs.slides[1].notes_slide.notes_text_frame.text
    title = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text == "样式与内容分离")
    font = title.text_frame.paragraphs[0].runs[0].font
    assert font.name == "微软雅黑" and font.size.pt == pytest.approx(32, abs=0.01)
    with ZipFile(result.output) as output:
        for src, dst in result.metadata["native_assets"].items():
            assert output.read(dst) == original.parts[src]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    assert not (run / ".workspace/pptx-style-compile").exists()
    assert load_style(run, str(path.relative_to(run))).name == "brand"


def test_inspection_follows_presentation_order_and_lists_typography(tmp_path):
    source, _, _ = fixture(tmp_path)
    prs = Presentation(source)
    second = prs.slides.add_slide(prs.slide_layouts[6])
    second.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "FIRST VISIBLE"
    ids = prs.slides._sldIdLst
    ids.insert(0, ids[-1])
    prs.save(source)
    found = inspect_pptx(source, slide=1)
    assert found["slide_count"] == 2
    assert found["slides"][0]["layers"]["slide"]["shapes"][0]["text"] == "FIRST VISIBLE"
    assert inspect_pptx(source, slide=2)["slides"][0]["layers"]["slide"]["shapes"][1][
        "sizes_pt"
    ] == [37]


@pytest.mark.parametrize("field", ["density", "slots", "layout", "page_role", "structure"])
def test_appearance_profile_rejects_style_of_organization_fields(tmp_path, field):
    _, run, path = fixture(tmp_path)
    data = yaml.safe_load(path.read_text())
    data[field] = "inappropriate"
    path.write_text(yaml.safe_dump(data))
    with pytest.raises(StyleError, match="density or slots"):
        render_deck(plan(), run, run / "bad.pptx")
    assert not (run / "bad.pptx").exists()


def test_unbound_and_stale_styles_do_not_fall_back(tmp_path):
    source, run, path = fixture(tmp_path)
    invalid = plan()
    invalid["deck"]["style"]["pptx_style"] = ""
    with pytest.raises(StyleError):
        render_deck(invalid, run, run / "invalid.pptx")
    draft = import_style(source, "draft", run)
    with pytest.raises(StyleError, match="surfaces"):
        load_style(run, str(draft.relative_to(run)))
    with pytest.raises(StyleError, match="templates/"):
        load_style(run, "../brand.pptx")
    data = yaml.safe_load(path.read_text())
    data["surfaces"]["content"]["keep"]["slide"] = [999]
    data["surfaces"]["content"].pop("protect")
    path.write_text(yaml.safe_dump(data))
    with pytest.raises(StyleError, match="999"):
        load_style(run, str(path.relative_to(run)))
    (path.parent / "source.pptx").write_bytes(b"changed")
    with pytest.raises(StyleError, match="hash changed"):
        load_style(run, str(path.relative_to(run)))


def test_content_overflow_fails_without_replacing_previous_output(tmp_path):
    _, run, _ = fixture(tmp_path)
    output = run / "out.pptx"
    render_deck(plan(), run, output)
    previous = output.read_bytes()
    bad = plan()
    bad["deck"]["pages"][0]["title"] = "标题过长" * 100
    with pytest.raises(StyleError, match="does not fit"):
        render_deck(bad, run, output)
    assert output.read_bytes() == previous
    assert not (run / ".workspace/pptx-style-compile").exists()


def test_charts_remain_editable_and_bound_to_evidence(tmp_path):
    _, run, _ = fixture(tmp_path)
    p = plan()
    p["deck"]["pages"][0]["visual"] = {
        "chart": {
            "type": "column",
            "series": [
                {"label": "before", "value_from": "E001"},
                {"label": "after", "value_from": "E002"},
            ],
        }
    }
    evidence = {
        "items": [
            {"id": "E001", "value": {"number": 20, "unit": "ms"}},
            {"id": "E002", "value": {"number": 10, "unit": "ms"}},
        ]
    }
    result = render_deck(p, run, run / "chart.pptx", evidence=evidence)
    prs = Presentation(result.output)
    chart = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values) == [20, 10]
    with ZipFile(result.output) as archive:
        assert any(n.startswith("ppt/embeddings/") for n in archive.namelist())


def test_page_edits_leave_other_content_and_style_parts_unchanged(tmp_path):
    _, run, _ = fixture(tmp_path)
    p = plan()
    first = render_deck(p, run, run / "first.pptx")
    revised = deepcopy(p)
    revised["deck"]["pages"][1]["title"] = "只改本页"
    second = render_deck(revised, run, run / "second.pptx")
    with ZipFile(first.output) as a, ZipFile(second.output) as b:
        assert a.read("ppt/slides/slide1.xml") == b.read("ppt/slides/slide1.xml")
        for name in a.namelist():
            if "comh-style-" in name:
                assert a.read(name) == b.read(name)


def test_html_and_conflicting_backgrounds_fail_explicitly(tmp_path):
    _, run, _ = fixture(tmp_path)
    with pytest.raises(StyleError, match="HTML"):
        render_html_deck(plan(), run, run / "wrong.html")
    p = plan()
    p["deck"]["pages"][0]["visual"] = {"background": {"asset": "x.png"}}
    with pytest.raises(StyleError, match="conflicts"):
        render_deck(p, run, run / "bad.pptx")


def test_protected_logo_collision_is_a_compile_error(tmp_path):
    _, run, path = fixture(tmp_path)
    source = path.parent / "source.pptx"
    prs = Presentation(source)
    logo = prs.slides[0].shapes[0]
    logo.left, logo.top, logo.width, logo.height = Inches(1), Inches(0.9), Inches(8), Inches(1)
    prs.save(source)
    data = yaml.safe_load(path.read_text())
    data["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    path.write_text(yaml.safe_dump(data))
    with pytest.raises(StyleError, match="overlaps protected"):
        render_deck(plan(), run, run / "bad.pptx")


def test_cli_inspect_and_existing_dialogue_guards(tmp_path, capsys):
    source, _, _ = fixture(tmp_path)
    assert main(["style-inspect", str(source), "--slide", "1"]) == 0
    assert "Appearance only" in capsys.readouterr().out
    root = tmp_path / "guarded"
    init_run(root)
    assert main(["style-import", str(source), "--name", "brand", "--run", str(root)]) == 0
    assert main(["render", "deck", "--run", str(root)]) == 2
    assert not (root / "build/deck.pptx").exists()


def test_native_background_is_checked_instead_of_palette_declaration(tmp_path):
    from comh.pptx_style.compiler import resolve_plan

    _, run, path = fixture(tmp_path)
    resolve_plan(plan(), run, allow_dark=False)
    source = path.parent / "source.pptx"
    prs = Presentation(source)
    prs.slide_master.background.fill.fore_color.rgb = RGBColor.from_string("000000")
    prs.save(source)
    data = yaml.safe_load(path.read_text())
    data["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    assert data["tokens"]["colors"]["background"] == "FFFFFF"
    path.write_text(yaml.safe_dump(data))
    with pytest.raises(StyleError, match="light-background"):
        resolve_plan(plan(), run, allow_dark=False)


def test_nested_icon_selection_does_not_import_sibling_content(tmp_path):
    from comh.pptx_style.package import NS

    _, run, path = fixture(tmp_path)
    source = path.parent / "source.pptx"
    prs = Presentation(source)
    slide = prs.slides[0]
    group = slide.shapes.add_group_shape()
    icon = group.shapes.add_picture(str(tmp_path / "logo.png"), 0, 0, Inches(0.2), Inches(0.2))
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(2)).text = "PRIVATE SIBLING"
    prs.save(source)
    data = yaml.safe_load(path.read_text())
    data["sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    data["surfaces"]["content"]["keep"]["slide"] = [icon.shape_id]
    data["surfaces"]["content"].pop("protect")
    path.write_text(yaml.safe_dump(data))
    result = render_deck(plan(), run, run / "out.pptx")
    package = Package(result.output)
    slide = package.tree(package.slides()[0])
    assert slide.find(".//p:grpSp/p:pic", NS) is not None
    assert "PRIVATE SIBLING" not in "".join(slide.xpath(".//a:t/text()", namespaces=NS))


def test_user_content_icons_survive_style_composition(tmp_path, monkeypatch):
    from comh.render import icons

    _, run, _ = fixture(tmp_path)
    monkeypatch.setattr(icons, "icon_png", lambda *a, **kw: tmp_path / "logo.png")
    p = plan()
    p["deck"]["pages"][0]["support_points"] = [{"point": "保留有意义的图标", "icon": "check"}]
    result = render_deck(p, run, run / "out.pptx")
    prs = Presentation(result.output)
    assert any(s.name == "icon:check" for s in prs.slides[0].shapes)


def test_failed_compile_records_failure_and_invalidates_old_build(tmp_path):
    from test_closed_loop import MIN_DECK_PLAN, _pipeline_through_narrative, _write
    from workflow_helpers import approve

    from comh.manifest import Manifest, RunError

    source, _, profile = fixture(tmp_path)
    root = init_run(tmp_path / "guarded")
    bound = import_style(source, "brand", root)
    bound.write_bytes(profile.read_bytes())
    _pipeline_through_narrative(root, appearance={"selection": "delegated", "review": "delegated"})
    p = yaml.safe_load(MIN_DECK_PLAN)
    p["deck"]["style"] = {"pptx_style": "templates/brand/style.yaml"}
    _write(root, "projection/deck_plan.yaml", yaml.safe_dump(p))
    assert main(["save", "deck_plan", "--run", str(root)]) == 0
    assert approve(root, "deck_outline") == 0
    assert main(["render", "deck", "--run", str(root)]) == 0
    manifest = Manifest.load(root)
    manifest.require_build("deck")
    output = manifest.output_path("deck")
    previous = output.read_bytes()
    data = yaml.safe_load(bound.read_text())
    data["tokens"]["sizes"]["content_title"] = 96
    bound.write_text(yaml.safe_dump(data))
    with pytest.raises(RunError, match="stale"):
        Manifest.load(root).require_build("deck")
    assert main(["render", "deck", "--run", str(root)]) == 2
    assert output.read_bytes() == previous
    with pytest.raises(RunError, match="no successful build"):
        Manifest.load(root).require_build("deck")
    report = yaml.safe_load((root / "qa/render-deck.yaml").read_text())
    assert report["published"] is False and report["count"]["error"] == 1
