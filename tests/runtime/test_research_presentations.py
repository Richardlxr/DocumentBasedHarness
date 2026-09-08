"""User-facing regression: edit, save, reopen; don't just count generated shapes."""

from copy import deepcopy
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from comh.presentation_profile import validate_density
from comh.render.deck import render_deck
from comh.render.editability import audit_editability
from comh.render.html_deck import render_html_deck
from comh.render.native_diagram import diagram_scene, diagram_svg
from comh.validate import validate_visuals
from comh.visible_text import page_texts


def specimen():
    return {
        "version": 1,
        "deck": {
            "title": "研究展示回归样张",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "硬件描述与运行接口的分离",
                    "visual": {
                        "arrangement": "full",
                        "diagram": {
                            "mermaid": "flowchart LR\n a[硬件描述] -->|参数| b[运行接口]\n"
                            " b --> c[应用程序]"
                        },
                    },
                },
                {
                    "id": "P02",
                    "page_role": "content",
                    "title": "配置方式与实验观测对照",
                    "visual": {
                        "table": {
                            "columns": ["方案", "配置方式", "观测耗时"],
                            "column_weights": [1, 2, 1],
                            "rows": [
                                [
                                    "基线",
                                    {"text": "配置随程序维护", "evidence": "E001"},
                                    {"value_from": "E001"},
                                ],
                                ["适配方案", "硬件描述独立维护", {"value_from": "E002"}],
                            ],
                        }
                    },
                },
                {
                    "id": "P03",
                    "page_role": "content",
                    "title": "从功能定位到跨核配置机制",
                    "visual": {"arrangement": "columns"},
                    "support_points": [
                        {"point": "功能定位", "detail": "说明控制对象与已有技术的边界。"},
                        {"point": "裸核适配", "detail": "将原有平台能力接入裸核运行环境。"},
                        {"point": "跨核配置机制", "detail": "解释配置必须由所属核心执行的原因。"},
                    ],
                },
            ],
        },
    }


EVIDENCE = {
    "items": [
        {"id": "E001", "value": {"number": 12, "unit": "ms"}},
        {"id": "E002", "value": {"number": 8, "unit": "ms"}},
    ]
}


def test_editable_diagram_and_table_survive_save_reopen(tmp_path):
    plan = specimen()
    plan["deck"]["style"] = {"animations": True}
    plan["deck"]["pages"][0]["reveal"] = [{"elements": ["visual"], "verb": "appear"}]
    result = render_deck(plan, tmp_path, tmp_path / "native.pptx", language="zh", evidence=EVIDENCE)
    prs = Presentation(result.output)
    assert not result.findings
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for s in prs.slides for shape in s.shapes)
    nodes = [s for s in prs.slides[0].shapes if s.name.startswith("diagram-node:")]
    nodes[0].text_frame.paragraphs[0].runs[0].text = "修改后的硬件描述"
    table = next(s.table for s in prs.slides[1].shapes if s.has_table)
    assert table.cell(1, 2).text == "12 ms"
    table.cell(1, 2).text = "15 ms"
    prs.save(tmp_path / "edited.pptx")
    edited = Presentation(tmp_path / "edited.pptx")
    assert any("修改后的" in s.text for s in edited.slides[0].shapes if s.has_text_frame)
    assert next(s.table for s in edited.slides[1].shapes if s.has_table).cell(1, 2).text == "15 ms"
    ids = result.shape_map["P01"]["visual"]
    targets = {int(e.get("spid")) for e in prs.slides[0]._element.iter(qn("p:spTgt"))}
    assert set(ids) <= targets  # all nodes, segments and labels reveal together
    assert result.metadata["editability"]["pages"][1]["tables"] == 1


def test_html_uses_same_labels_numbers_and_searchable_svg(tmp_path):
    result = render_html_deck(specimen(), tmp_path, tmp_path / "deck.html", evidence=EVIDENCE)
    text = result.output.read_text()
    assert "<svg" in text and "<text " in text and "参数" in text
    assert "12 ms" in text and "8 ms" in text and '<th scope="col">' in text
    assert 'class="reading-columns"' in text
    assert '<img src="data:image/png;base64' not in text


@pytest.mark.parametrize(
    "change,expected",
    [
        (lambda v: v["table"]["rows"][0].pop(), "column count"),
        (lambda v: v["table"].update(column_weights=[1, 0, 1]), "column_weights"),
        (lambda v: v.update(diagram={"mermaid": "flowchart LR\n a --> b"}), "carrier"),
        (lambda v: v["table"]["rows"][0][2].update(value_from="E999"), "missing evidence"),
    ],
)
def test_invalid_tables_and_competing_carriers_fail_before_render(change, expected):
    plan = specimen()
    change(plan["deck"]["pages"][1]["visual"])
    findings = validate_visuals({"deck_plan": plan, "evidence": EVIDENCE})
    assert any(f.severity == "error" and expected in f.detail for f in findings)


def test_table_copy_is_visible_to_audience_checks():
    page = specimen()["deck"]["pages"][1]
    texts = page_texts(page)
    assert ("visual.table.rows[0][1]", "配置随程序维护") in texts
    assert ("visual.table.columns[2]", "观测耗时") in texts


def test_oversized_table_fails_without_font_shrink(tmp_path):
    plan = specimen()
    plan["deck"]["pages"] = [plan["deck"]["pages"][1]]
    plan["deck"]["pages"][0]["visual"]["table"]["rows"] *= 20
    with pytest.raises(ValueError, match="table does not fit"):
        render_deck(plan, tmp_path, tmp_path / "too-many.pptx", evidence=EVIDENCE)


def test_branch_routes_diamond_and_cylinder_preserve_semantics(tmp_path):
    source = "flowchart TB\n a{选择} -->|是| b[(数据)]\n a -->|否| c[退出]"
    scene = diagram_scene(source)
    assert any(n["style"].get("shape") == "rhombus" for n in scene.nodes)
    assert "<polygon" in diagram_svg(scene)
    plan = specimen()
    plan["deck"]["pages"] = [plan["deck"]["pages"][0]]
    plan["deck"]["pages"][0]["visual"]["diagram"]["mermaid"] = source
    result = render_deck(plan, tmp_path, tmp_path / "branch.pptx")
    slide = Presentation(result.output).slides[0]
    shapes = [s for s in slide.shapes if s.name.startswith("diagram-node:")]
    assert {MSO_AUTO_SHAPE_TYPE.DIAMOND, MSO_AUTO_SHAPE_TYPE.CAN} <= {
        s.auto_shape_type for s in shapes
    }
    assert any(s._element.find(".//" + qn("a:tailEnd")) is not None for s in slide.shapes)


def test_text_streak_is_advisory_and_draft_is_exempt():
    plan = specimen()
    page = plan["deck"]["pages"][2]
    page["visual"] = {}
    plan["deck"]["pages"] = [{**deepcopy(page), "id": f"P{i}"} for i in range(3)]
    artifacts = {
        "brief": {
            "media": [{"surface": "presentation"}],
            "language": "zh",
            "presentation": {"setting": "academic"},
        },
        "deck_plan": plan,
    }
    assert any(
        f.check == "presentation:text-only-sequence" and f.severity == "warn"
        for f in validate_density(artifacts)
    )
    plan["draft"] = True
    assert not validate_density(artifacts)


def test_raster_dominance_is_review_signal_for_source_images(tmp_path):
    from PIL import Image
    from pptx.util import Inches

    path = tmp_path / "source.png"
    Image.new("RGB", (100, 100)).save(path)
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(
        str(path), 0, 0, prs.slide_width, prs.slide_height
    )
    inventory, findings = audit_editability(prs, ["P01"])
    assert inventory["pages"][0]["pictures"] == 1
    assert findings[0].severity == "warn"
    prs.slides[0].shapes[0].width = Inches(1)
    assert not audit_editability(prs, ["P01"])[1]


def test_tiny_diagram_fails_instead_of_rasterizing(tmp_path):
    from comh.render.native_diagram import render_native_diagram
    from comh.render.theme import render_theme

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    with pytest.raises(ValueError, match="below 17pt"):
        render_native_diagram(
            slide,
            diagram_scene("flowchart LR\n a --> b"),
            render_theme(None, "zh"),
            x=1,
            y=1,
            w=1,
            h=1,
        )


def test_dark_theme_diagram_labels_keep_contrasting_background(tmp_path):
    plan = specimen()
    plan["deck"]["pages"] = plan["deck"]["pages"][:1]
    plan["deck"]["style"] = {"template": "midnight"}
    result = render_deck(plan, tmp_path, tmp_path / "dark.pptx", language="zh")
    slide = Presentation(result.output).slides[0]
    label = next(s for s in slide.shapes if s.name.startswith("diagram-label:"))
    assert str(label.fill.fore_color.rgb) == "FFFFFF"
    svg = diagram_svg(diagram_scene("flowchart LR\n a[A] -->|X| b[B]", 'A"B'))
    assert 'font-family="A&quot;B"' in svg


def test_chart_titles_are_reader_visible_not_metadata():
    page = {"visual": {"chart": {"title": "Reader-visible heading"}}}
    assert ("visual.chart.title", "Reader-visible heading") in page_texts(page)


if __name__ == "__main__":
    root = Path(__import__("sys").argv[1])
    root.mkdir(parents=True, exist_ok=True)
    render_deck(specimen(), root, root / "specimen.pptx", language="zh", evidence=EVIDENCE)
    render_html_deck(specimen(), root, root / "specimen.html", language="zh", evidence=EVIDENCE)


def test_bidirectional_flow_retains_both_arrowheads(tmp_path):
    plan = specimen()
    plan["deck"]["pages"] = plan["deck"]["pages"][:1]
    source = "flowchart LR\n a[Sender] <--> b[Receiver]"
    plan["deck"]["pages"][0]["visual"]["diagram"]["mermaid"] = source
    result = render_deck(plan, tmp_path, tmp_path / "two-way.pptx")
    slide = Presentation(result.output).slides[0]
    assert any(s._element.find(".//" + qn("a:headEnd")) is not None for s in slide.shapes)
    assert any(s._element.find(".//" + qn("a:tailEnd")) is not None for s in slide.shapes)
    assert diagram_svg(diagram_scene(source)).count("<polygon") == 2
