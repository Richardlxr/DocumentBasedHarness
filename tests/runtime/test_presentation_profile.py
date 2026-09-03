from copy import deepcopy

from pptx import Presentation
from test_dialogue import cli, until_brief, until_outline, write
from test_workflow_integrity import BRIEF, DECK
from workflow_helpers import aligned, approve

from comh.cli import main
from comh.context import context_pack
from comh.decision_checks import brief_errors
from comh.manifest import Manifest
from comh.presentation_profile import policy_errors, resolve_profile, validate_density
from comh.render.deck import render_deck
from comh.render.metrics import fixed_text_sizes
from comh.visible_text import page_texts


def academic():
    brief = deepcopy(BRIEF)
    brief["presentation"] = {"setting": "academic"}
    return aligned(brief)


def test_academic_default_explicit_override_and_general_are_distinct():
    brief = academic()
    assert resolve_profile(brief)["profile"] == "academic-rich"
    brief["presentation"]["profile"] = "concise"
    assert resolve_profile(brief)["text_budget"]["zh_chars"] == [60, 180]
    brief["presentation"] = {"setting": "general"}
    assert resolve_profile(brief)["profile"] == "balanced"
    brief["presentation"]["profile"] = "academic-rich"
    assert resolve_profile(brief)["profile"] == "academic-rich"


def test_overrides_do_not_mutate_other_runs_or_interpret_open_dimensions():
    brief = academic()
    brief["presentation"]["overrides"] = {
        "text_budget": {"zh_chars": [200, 400]},
        "layout_patterns": ["按用户指定的两块结构组织"],
    }
    resolved = resolve_profile(brief)
    assert resolved["text_budget"]["zh_chars"] == [200, 400]
    assert resolve_profile(academic())["text_budget"]["zh_chars"] == [300, 550]
    brief["spec"] = {"dimensions": {"evidence_density": "unusual-user-vocabulary"}}
    assert resolve_profile(brief) == resolved
    assert not policy_errors(brief)
    brief["presentation"]["overrides"]["text_budget"]["zh_chars"] = [400, 200]
    assert policy_errors(brief)


def test_intent_is_aligned_at_gate_one_without_an_extra_approval_round(tmp_path):
    brief = academic()
    brief["alignment"]["presentation"].update(
        source="default", status="proposed", basis="学术汇报采用充实预设"
    )
    root = until_brief(tmp_path, brief)
    assert approve(root, "brief") == 0
    brief["presentation"] = {"setting": "academic", "profile": "concise"}
    assert brief_errors(brief)  # Actual value and accepted proposed value must agree.
    no_policy = academic()
    del no_policy["presentation"], no_policy["alignment"]["presentation"]
    assert any("presentation" in e for e in brief_errors(no_policy))
    no_policy["media"] = [{"medium": "markdown", "surface": "report"}]
    no_policy["alignment"]["media"]["value"] = no_policy["media"]
    assert not any("presentation" in e for e in brief_errors(no_policy))


def test_profile_is_progressively_disclosed_and_available_without_run(tmp_path, capsys):
    root = until_outline(tmp_path, academic())
    manifest = Manifest.load(root)
    assert "presentation_profile" not in context_pack(manifest, stage="narrative")
    pack = context_pack(manifest, stage="deck", node="P01")
    assert pack["presentation_profile"]["profile"] == "academic-rich"
    assert pack["presentation_profile"]["layout_patterns"]
    assert pack["copy_contract"]["audience"] == academic()["audience"]
    assert main(["presentation-profiles", "academic-rich"]) == 0
    assert "学术充实" in capsys.readouterr().out
    assert main(["presentation-profiles", "missing"]) == 2


def test_density_is_advisory_does_not_force_text_onto_charts_or_outline():
    plan = deepcopy(DECK)
    plan["deck"]["pages"][0]["support_points"] = ["第一点", "第二点", "第三点"]
    artifacts = {"brief": academic(), "deck_plan": plan}
    assert any(
        f.check == "density:sparse" and f.severity == "warn" for f in validate_density(artifacts)
    )
    plan["deck"]["pages"][0]["visual"] = {"chart": {"type": "column", "series": []}}
    assert not validate_density(artifacts)
    plan["deck"]["pages"][0]["visual"] = {}
    plan["deck"]["pages"][0]["page_role"] = "cover"
    assert not validate_density(artifacts)
    plan["deck"]["pages"][0]["page_role"] = "content"
    plan["draft"] = True
    assert not validate_density(artifacts)
    del plan["draft"]
    plan["deck"]["pages"][0]["support_points"] = ["内容" * 500]
    assert any(f.check == "density:dense" for f in validate_density(artifacts))


def test_sparse_advice_does_not_block_formal_render(tmp_path):
    root = until_outline(tmp_path, academic())
    assert approve(root, "deck_outline") == 0
    assert cli(root, "render", "deck") == 0
    changed = academic()
    changed["presentation"]["overrides"] = {"text_budget": {"zh_chars": [900, 100]}}
    changed = aligned(changed)
    write(root, "brief/brief.yaml", changed)
    assert cli(root, "save", "brief") == 0  # Can save a draft, cannot accept it.
    assert approve(root, "brief") == 2


def test_rich_page_fits_existing_renderer_without_shrinking_type(tmp_path):
    # A representative text-heavy page, not a claim about every layout or player.
    plan = deepcopy(DECK)
    page = plan["deck"]["pages"][0]
    page["title"] = "测量结果应同时交代实验条件与可支持的判断"
    page["support_points"] = [
        {
            "point": "记录实验设置，让读者知道比较是否成立",
            "detail": (
                "比较前先说明运行环境、输入数据、负载方式和采样过程，再列出基线与改动后的设置。"
                "如果两个实验除了目标因素之外还存在差异，应在结果旁说明差异可能造成的影响。"
                "原始记录需要保留测量位置、重复运行方式和统计口径，便于检查异常值与缺失数据。"
                "图中的每组结果应能回到对应实验，而不是只留下一个没有出处的平均值。"
                "尚未完成的校准与未覆盖的负载应单独列出，不能把没有观察到问题解释为问题不存在。"
            ),
        },
        {
            "point": "解释观察结果，同时说明尚未验证的部分",
            "detail": (
                "解释结果时，先描述基线与改动后出现了什么差别，再讨论哪些机制可能解释这一差别。"
                "观察、推断和后续假设需要分别表达，避免把一种合理解释写成已经证实的唯一原因。"
                "相同设置下重复得到相同输出可以支持实验可复现，但不能单独证明模型与实际设备一致。"
                "实际应用还需要检查真实负载、运行环境与测量方法的差异，并安排对应的对照验证。"
                "如果现有记录不足以回答某个问题，应直接保留问题和需要补充的证据，而不是用笼统评价结束讨论。"
            ),
        },
    ]
    count = len("".join(t for _, t in page_texts(page)))
    assert 400 <= count <= 550
    assert not validate_density({"brief": academic(), "deck_plan": plan})
    with fixed_text_sizes():
        result = render_deck(plan, tmp_path, tmp_path / "rich.pptx", language="zh-CN")
    assert not [f for f in result.findings if f.verdict == "fail"], result.findings
    slide = Presentation(result.output).slides[0]
    sizes = {
        paragraph.text: paragraph.runs[0].font.size.pt
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        if paragraph.runs
    }
    for entry in page["support_points"]:
        assert sizes[entry["point"]] == 24
        assert sizes[entry["detail"]] == 17
    cards = [shape for shape in slide.shapes if shape.name.startswith("card:point[")]
    assert max((shape.top + shape.height) / 914400 for shape in cards) <= 6.8

    # A real conclusion still owns its band; density must not suppress overflow.
    page["callout"] = {"text": "在给定条件下解释结果，并列出需要补充的验证。"}
    with fixed_text_sizes():
        cramped = render_deck(plan, tmp_path, tmp_path / "cramped.pptx", language="zh-CN")
    assert any(f.check == "layout" and f.verdict == "fail" for f in cramped.findings)
