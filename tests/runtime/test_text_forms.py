"""Text body forms and semantic tables: relation-named layouts, measured fit,
the same copy and addresses on both surfaces, and advisory variety checks."""

from copy import deepcopy

import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

from comh.dialogue import _page_carrier
from comh.presentation_profile import (
    presentation_structure_errors,
    validate_layout_variety,
    validate_packing,
)
from comh.render.deck import render_deck
from comh.render.html_deck import render_html_deck
from comh.render.text_forms import FormFitError, form_errors
from comh.validate import validate_layout_fit, validate_visuals
from comh.visible_text import page_texts

EVIDENCE = {
    "items": [
        {"id": "E001", "value": {"number": 220, "unit": "ms"}},
        {"id": "E002", "value": {"number": 180, "unit": "ms"}},
    ]
}
BRIEF = {
    "media": [{"surface": "presentation"}],
    "language": "zh",
    "presentation": {"setting": "academic"},
}


def pair(point, detail, **extra):
    return {"point": point, "detail": detail, **extra}


FORM_PAGES = {
    "steps": [pair(f"第{i}步做什么", "这一步的产出可以检查。") for i in range(1, 5)],
    "grid": [pair(name, "这个问题检验一项具体能力。") for name in "甲乙丙丁"],
    "statement": [
        pair("分离描述后冷启动没有变长", "差异在测量噪声内。"),
        pair("测量条件", "单板、室温。"),
        pair("未覆盖", "多核同时启动。"),
    ],
    "qa": [pair("生成代码会更慢吗？", "测量上看不出差别。"), pair("写错怎么办？", "编译前报错。")],
    "definition": [pair("裸核", "不运行操作系统的核心。"), pair("跨核配置", "修改他核外设。")],
    "status": [
        pair("初始化代码生成", "两块板卡跑通。", status="已验证", tone="positive"),
        pair("跨核配置", "多核待测。", status="部分验证", tone="partial"),
        pair("第三方芯片", "尚未实现。", status="设计目标", tone="pending"),
    ],
}


def form_plan(form, points=None, **page):
    return {
        "version": 1,
        "deck": {
            "title": "样张",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "一页文字形态",
                    "visual": {"arrangement": form},
                    "support_points": deepcopy(points or FORM_PAGES[form]),
                    **page,
                }
            ],
        },
    }


@pytest.mark.parametrize("form", sorted(FORM_PAGES))
def test_each_form_renders_editable_text_on_both_surfaces(tmp_path, form):
    plan = form_plan(form, callout={"text": "结论条保持在正文下方。"})
    plan["deck"]["style"] = {"animations": True}
    plan["deck"]["pages"][0]["reveal"] = [
        {"elements": ["support_points[0]"], "verb": "fade_in", "trigger": "click"}
    ]
    result = render_deck(plan, tmp_path, tmp_path / "f.pptx", language="zh")
    assert not result.findings, [f.detail for f in result.findings]
    slide = Presentation(result.output).slides[0]
    text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    for entry in FORM_PAGES[form]:
        assert entry["point"] in text and entry["detail"] in text
        if "status" in entry:
            assert entry["status"] in text
    shapes = result.shape_map["P01"]
    points = len(FORM_PAGES[form])
    assert all(shapes[f"support_points[{i}]"] for i in range(points))
    targets = {int(e.get("spid")) for e in slide._element.iter(qn("p:spTgt"))}
    assert set(shapes["support_points[0]"]) <= targets

    html = render_html_deck(plan, tmp_path, tmp_path / "f.html", evidence=None)
    page = html.output.read_text()
    assert f"form form-{form}" in page
    assert 'class="form-item form-' in page and "data-fragment-index" in page
    assert FORM_PAGES[form][-1]["detail"] in page


def test_steps_badges_number_the_sequence_and_details_share_a_baseline(tmp_path):
    points = deepcopy(FORM_PAGES["steps"])
    points[0]["point"] = "第一步的说明写得比其他步骤更长，需要换到第二行甚至第三行显示"
    result = render_deck(form_plan("steps", points), tmp_path, tmp_path / "s.pptx", language="zh")
    slide = Presentation(result.output).slides[0]
    badges = [s for s in slide.shapes if s.name.startswith("form:oval")]
    assert [b.text_frame.text for b in badges] == ["1", "2", "3", "4"]
    details = [s for s in slide.shapes if s.has_text_frame and "产出可以检查" in s.text_frame.text]
    assert len({d.top for d in details}) == 1


def test_five_steps_use_the_vertical_ladder(tmp_path):
    points = [pair(f"步骤{i}", "说明。") for i in range(5)]
    result = render_deck(form_plan("steps", points), tmp_path, tmp_path / "v.pptx", language="zh")
    badges = [
        s for s in Presentation(result.output).slides[0].shapes if s.name.startswith("form:oval")
    ]
    assert len({b.left for b in badges}) == 1 and len({b.top for b in badges}) == 5


@pytest.mark.parametrize(
    "form,change,expected",
    [
        ("grid", lambda p: p.append(pair("戊", "多出一项")), "four or six"),
        ("qa", lambda p: p[0].pop("detail"), "needs a detail"),
        ("status", lambda p: p[0].pop("status"), "visible status label"),
        ("status", lambda p: p[0].update(tone="great"), "tone 'great'"),
        ("definition", lambda p: p[0].update(icon="bolt"), "icon does not render"),
        ("steps", lambda p: p.extend([pair("x", "y")] * 3), "requires 2–5"),
    ],
)
def test_form_structure_errors_are_source_oriented(form, change, expected):
    plan = form_plan(form)
    change(plan["deck"]["pages"][0]["support_points"])
    errors = presentation_structure_errors(plan["deck"]["pages"][0])
    assert any(expected in e for e in errors), errors
    findings = validate_visuals({"deck_plan": plan, "evidence": EVIDENCE})
    assert any(f.severity == "error" and expected in f.detail for f in findings)


def test_forms_refuse_competing_carriers_and_status_outside_status_form():
    page = form_plan("qa")["deck"]["pages"][0]
    page["visual"]["chart"] = {"series": [{"value_from": "E001"}]}
    assert any("text body form" in e for e in form_errors(page))
    listed = {"support_points": [pair("a", "b", status="已验证")], "visual": {}}
    assert any("status arrangement" in e for e in form_errors(listed))


def test_overflowing_form_fails_at_validate_and_render_without_shrinking(tmp_path):
    long = "这一段解释写得非常长，" * 30
    plan = form_plan("qa", [pair("问题一？", long), pair("问题二？", long)])
    findings = validate_layout_fit({"deck_plan": plan, "evidence": EVIDENCE}, tmp_path)
    assert any(f.check == "layout-fit" and f.severity == "error" for f in findings)
    with pytest.raises(FormFitError, match="does not fit at"):
        render_deck(plan, tmp_path, tmp_path / "x.pptx", language="zh")


def test_layout_dry_run_is_skipped_out_loud_for_pptx_style():
    plan = form_plan("qa")
    plan["deck"]["style"] = {"pptx_style": "templates/x"}
    findings = validate_layout_fit({"deck_plan": plan}, None)
    assert findings and findings[0].severity == "info" and "skipped" in findings[0].detail


# -- tables -------------------------------------------------------------------


def table_plan(**table):
    spec = {
        "columns": ["方案", "配置", "冷启动", "功能测试"],
        "header_column": True,
        "highlight": {"rows": [1]},
        "note": "冷启动为 20 次测量中位数。",
        "rows": [
            ["基线", "随程序维护", {"value_from": "E001"}, {"status": "通过", "tone": "positive"}],
            ["适配", "独立描述", {"value_from": "E002"}, {"status": "通过", "tone": "positive"}],
            ["第二板卡", "同一描述", {"missing": "未测"}, {"status": "进行中", "tone": "pending"}],
        ],
        **table,
    }
    return {
        "version": 1,
        "deck": {
            "title": "表格",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "content",
                    "title": "适配方案与基线相当",
                    "visual": {"arrangement": "full", "table": spec},
                }
            ],
        },
    }


def test_semantic_table_marks_headers_numbers_highlights_status_and_missing(tmp_path):
    result = render_deck(table_plan(), tmp_path, tmp_path / "t.pptx", evidence=EVIDENCE)
    assert not result.findings, [f.detail for f in result.findings]
    slide = Presentation(result.output).slides[0]
    table = next(s.table for s in slide.shapes if s.has_table)
    assert table.cell(1, 2).text == "220 ms"
    assert table.cell(1, 2).text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    assert table.cell(0, 2).text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold  # row label
    assert table.cell(2, 1).text_frame.paragraphs[0].runs[0].font.bold  # highlighted row
    assert not table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.bold
    assert table.cell(3, 2).text == "未测"
    assert table.cell(3, 2).text_frame.paragraphs[0].runs[0].font.italic
    note = next(s for s in slide.shapes if s.name == "table-note")
    assert "中位数" in note.text_frame.text
    assert note.shape_id in result.shape_map["P01"]["visual"]
    # edited value survives save/reopen like any native table
    table.cell(1, 2).text = "230 ms"
    Presentation(result.output)

    page = render_html_deck(table_plan(), tmp_path, tmp_path / "t.html", evidence=EVIDENCE)
    html = page.output.read_text()
    assert '<th scope="row">基线</th>' in html
    assert 'class="num"' in html and 'class="hl' in html
    assert 'form-pill tone-pending">进行中' in html and '<em class="muted">未测</em>' in html
    assert 'class="table-note' in html


def test_cell_borders_precede_fill_in_cell_properties(tmp_path):
    result = render_deck(table_plan(), tmp_path, tmp_path / "b.pptx", evidence=EVIDENCE)
    table = next(s.table for s in Presentation(result.output).slides[0].shapes if s.has_table)
    tc_pr = table.cell(0, 0)._tc.tcPr
    tags = [child.tag.split("}")[1] for child in tc_pr]
    assert tags[:4] == ["lnL", "lnR", "lnT", "lnB"]
    header_bottom = tc_pr.find(qn("a:lnB"))
    assert header_bottom.get("w") == str(2 * 12700)


@pytest.mark.parametrize(
    "change,expected",
    [
        ({"highlight": {"rows": [9]}}, "highlight row 9"),
        ({"highlight": {"cells": [[0, 7]]}}, "highlight cell"),
        ({"rows": [["a", "b", {"missing": " "}, "d"]]}, "missing cells"),
        ({"rows": [["a", "b", "c", {"status": "ok", "tone": "loud"}]]}, "status tone"),
    ],
)
def test_table_annotations_are_validated(change, expected):
    plan = table_plan(**change)
    findings = validate_visuals({"deck_plan": plan, "evidence": EVIDENCE})
    assert any(f.severity == "error" and expected in f.detail for f in findings)


def test_table_annotations_are_audience_visible_copy():
    texts = dict(page_texts(table_plan()["deck"]["pages"][0]))
    assert texts["visual.table.note"].startswith("冷启动")
    assert texts["visual.table.rows[2][2]"] == "未测"
    assert texts["visual.table.rows[0][3]"] == "通过"
    status_page = form_plan("status")["deck"]["pages"][0]
    assert ("support_points[0].status", "已验证") in page_texts(status_page)


def test_oversized_table_with_note_fails_in_validate_dry_run(tmp_path):
    plan = table_plan()
    plan["deck"]["pages"][0]["visual"]["table"]["rows"] *= 12
    findings = validate_layout_fit({"deck_plan": plan, "evidence": EVIDENCE}, tmp_path)
    assert any("table does not fit" in f.detail for f in findings)


def test_versus_table_compares_dimensions_and_keeps_its_callout(tmp_path):
    plan = {
        "version": 1,
        "deck": {
            "title": "对照",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "versus",
                    "title": "两种方案逐项对照",
                    "callout": {"text": "适配方案把改动集中到一个文件。"},
                    "visual": {
                        "table": {
                            "columns": ["维度", "基线", "适配"],
                            "highlight": {"columns": [2]},
                            "rows": [
                                ["新增外设", "3 个文件", "1 个文件"],
                                ["换板卡", "重写", "替换"],
                            ],
                        }
                    },
                }
            ],
        },
    }
    result = render_deck(plan, tmp_path, tmp_path / "vs.pptx", language="zh")
    assert not result.findings, [f.detail for f in result.findings]
    slide = Presentation(result.output).slides[0]
    table = next(s.table for s in slide.shapes if s.has_table)
    assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold  # row labels by default
    assert result.shape_map["P01"]["callout"]
    html = render_html_deck(plan, tmp_path, tmp_path / "vs.html").output.read_text()
    assert '<th scope="row">新增外设</th>' in html and "callout" in html
    plan["deck"]["pages"][0]["support_points"] = ["多余的栏目"]
    assert any("versus table" in e for e in presentation_structure_errors(plan["deck"]["pages"][0]))


# -- variety and shape hints ---------------------------------------------------


def deck_of(*pages):
    return {
        "version": 1,
        "deck": {
            "title": "d",
            "pages": [
                {"id": f"P{i:02d}", "page_role": "content", "title": "t", **page}
                for i, page in enumerate(pages, 1)
            ],
        },
    }


LIST = {"support_points": [pair("a", "b"), pair("c", "d")]}


def test_repeated_and_dominant_layouts_are_advisory():
    plan = deck_of(LIST, LIST, LIST, {"visual": {"arrangement": "qa"}, **LIST}, LIST)
    findings = validate_layout_variety({"deck_plan": plan})
    checks = {f.check for f in findings}
    assert checks == {"presentation:repeated-layout", "presentation:layout-monotony"}
    assert all(f.severity == "warn" for f in findings)
    assert "P01, P02, P03" in next(
        f.detail for f in findings if f.check == "presentation:repeated-layout"
    )
    plan["draft"] = True
    assert not validate_layout_variety({"deck_plan": plan})


def test_varied_forms_do_not_trigger_variety_or_text_streak_warnings():
    forms = ["steps", "qa", "definition", "status", "grid"]
    plan = deck_of(
        *({"visual": {"arrangement": f}, "support_points": FORM_PAGES[f]} for f in forms)
    )
    assert not validate_layout_variety({"deck_plan": plan})
    from comh.presentation_profile import validate_density

    streak = [
        f
        for f in validate_density({"brief": BRIEF, "deck_plan": plan})
        if f.check == "presentation:text-only-sequence"
    ]
    assert not streak


def test_outline_carrier_names_the_text_form():
    assert _page_carrier({"visual": {"arrangement": "steps"}}) == "text:steps"
    assert _page_carrier({"support_points": ["a"]}) == "text"
    assert _page_carrier({"visual": {"table": {"columns": ["a"]}}}) == "table"
    hero = {"page_role": "hero_split", "visual": {"table": {"columns": ["a"]}}}
    assert _page_carrier(hero) == "table"


@pytest.mark.parametrize(
    "points,check",
    [
        (
            [pair(n, "优点：快；限制：贵") for n in ("甲", "乙", "丙")],
            "layout:suggest-table",
        ),
        ([pair(f"第{n}步：准备", "说明") for n in "一二三"], "layout:suggest-steps"),
        ([pair(f"{n}. 准备", "说明") for n in (1, 2, 3)], "layout:suggest-steps"),
        ([pair("为什么慢？", "因为锁。"), pair("怎么修？", "拆锁。")], "layout:suggest-qa"),
    ],
)
def test_list_items_with_visible_structure_get_a_form_hint(points, check):
    plan = deck_of({"support_points": points})
    findings = validate_packing({"brief": BRIEF, "deck_plan": plan})
    assert any(f.check == check and f.severity == "warn" for f in findings)
    plan["deck"]["pages"][0]["visual"] = {"arrangement": "definition"}
    assert not any(f.check == check for f in validate_packing({"brief": BRIEF, "deck_plan": plan}))


def test_word_markers_only_run_for_covered_languages():
    points = [pair(w, "detail") for w in ("First, prepare", "Then build", "Finally ship")]
    plan = deck_of({"support_points": points})
    english = validate_packing({"brief": {**BRIEF, "language": "en"}, "deck_plan": plan})
    assert any(f.check == "layout:suggest-steps" for f in english)
    other = validate_packing({"brief": {**BRIEF, "language": "fr"}, "deck_plan": plan})
    assert not any(f.check == "layout:suggest-steps" for f in other)
    from comh.locale_support import uncovered_packs

    assert "layout-shape" in uncovered_packs({"language": "fr"})
