"""Unit and integration tests for deck layout enhancements, backgrounds, and animations."""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches

from comh.render.deck import render_deck
from comh.render.html_deck import render_html_deck
from comh.render.ooxml import set_shape_translucent_fill
from comh.validate import validate_assets, validate_visuals


def test_ooxml_translucent_fill():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(4), Inches(3))

    set_shape_translucent_fill(shape, RGBColor(255, 0, 0), alpha=0.65)

    spPr = shape._element.spPr
    solidFill = spPr.find(qn("a:solidFill"))
    assert solidFill is not None
    srgbClr = solidFill.find(qn("a:srgbClr"))
    assert srgbClr is not None
    alpha_elem = srgbClr.find(qn("a:alpha"))
    assert alpha_elem is not None
    assert alpha_elem.get("val") == "65000"


def test_render_all_layouts_pptx_and_html(tmp_path: Path):
    # Setup test assets and manifest
    assets_dir = tmp_path / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    img_path = assets_dir / "bg.png"
    Image.new("RGB", (400, 300), color=(50, 50, 50)).save(img_path)

    manifest_yaml = """
assets:
  - file: assets/bg.png
    origin_url: https://example.com/bg.png
    license: CC-BY-4.0
    fetched_at: "2026-09-02T12:00:00Z"
"""
    (assets_dir / "manifest.yaml").write_text(manifest_yaml, encoding="utf-8")

    evidence = {
        "items": [
            {"id": "E01", "value": {"number": 18.2, "unit": "%"}},
            {"id": "E02", "value": {"number": 220, "unit": "ms"}},
        ]
    }

    plan = {
        "deck": {
            "title": "Layout Test Deck",
            "style": {"template": "slate-tech"},
            "pages": [
                {
                    "id": "P01",
                    "page_role": "hero_split",
                    "title": "Hero Split Layout Demonstration",
                    "kicker": "ARCHITECTURE",
                    "support_points": [
                        {
                            "point": "Decoupled Architecture",
                            "detail": "Clean separation of concerns across tiers.",
                        },
                        {
                            "point": "Real-time Streaming",
                            "detail": "Ultra low latency event processing pipeline.",
                        },
                    ],
                    "callout": {"text": "Scales seamlessly to 10M DAU with 99.99% availability."},
                    "visual": {
                        "asset_refs": [{"ref": "assets/bg.png", "caption": "System Diagram"}],
                    },
                },
                {
                    "id": "P02",
                    "page_role": "fullscreen_backdrop",
                    "title": "Fullscreen Backdrop Layout",
                    "kicker": "VISION",
                    "support_points": [
                        {"point": "Global Reach", "detail": "Deployed in 15 regions worldwide."},
                    ],
                    "callout": {"text": "Mission critical infrastructure ready for deployment."},
                    "visual": {
                        "background": {
                            "asset": "assets/bg.png",
                            "opacity": 0.2,
                            "overlay": "frosted-glass",
                        },
                    },
                },
                {
                    "id": "P03",
                    "page_role": "timeline",
                    "title": "Evolution Timeline",
                    "kicker": "ROADMAP",
                    "support_points": [
                        {"point": "Phase 1: Alpha", "detail": "Core pipeline verification."},
                        {"point": "Phase 2: Beta", "detail": "User group trials."},
                        {"point": "Phase 3: GA", "detail": "Full production rollout."},
                    ],
                },
                {
                    "id": "P04",
                    "page_role": "versus",
                    "title": "Comparison Analysis",
                    "kicker": "TRADE-OFFS",
                    "support_points": [
                        {
                            "point": "Baseline Latency: 450ms",
                            "detail": "Synchronous blocking calls.",
                        },
                        {
                            "point": "Optimized Latency: 95ms",
                            "detail": "Async non-blocking event loop.",
                        },
                    ],
                    "visual": {"columns": ["BASELINE", "OPTIMIZED"]},
                },
            ],
        }
    }

    # 1. Validate assets and visuals
    artifacts = {"deck_plan": plan, "evidence": evidence}
    asset_findings = validate_assets(artifacts, tmp_path)
    assert not any(f.severity == "error" for f in asset_findings)

    visual_findings = validate_visuals(artifacts)
    assert not any(f.severity == "error" for f in visual_findings)

    # 2. Render PPTX
    pptx_out = tmp_path / "deck.pptx"
    pptx_result = render_deck(plan, tmp_path, pptx_out, language="en", evidence=evidence)
    assert pptx_out.is_file()
    assert pptx_result.theme == "slate-tech"
    assert not any(f.severity == "error" for f in pptx_result.findings)
    # versus column labels come from the spec on the pptx surface too
    prs = Presentation(str(pptx_out))
    pptx_text = "\n".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "BASELINE" in pptx_text and "OPTIMIZED" in pptx_text
    assert "方案 A" not in pptx_text  # renderers never invent column labels

    # 3. Render HTML
    html_out = tmp_path / "deck.html"
    html_result = render_html_deck(plan, tmp_path, html_out, language="en", evidence=evidence)
    assert html_out.is_file()
    assert html_result.theme == "slate-tech"
    html_text = html_out.read_text(encoding="utf-8")
    assert "role-hero_split" in html_text
    assert "role-fullscreen_backdrop" in html_text
    assert "role-timeline" in html_text
    assert "role-versus" in html_text
    assert "runCountUp" in html_text
    assert "hero-split" in html_text
    assert "milestone-badge" in html_text
    # versus labels from visual.columns, not renderer copy
    assert ">BASELINE</div>" in html_text and ">OPTIMIZED</div>" in html_text
    assert "方案 A" not in html_text
    # frosted glass blurs the ::after scrim layer, never an inline section style
    assert "has-bg frosted" in html_text
    assert ".has-bg.frosted::after" in html_text
    assert not re.search(r'<section[^>]*style="[^"]*backdrop-filter', html_text)
    # bar growth has an actual trigger, not just a static transition rule
    assert "growBars" in html_text and "data-grown" in html_text


def test_versus_without_columns_degrades_to_bare_ab(tmp_path: Path):
    plan = {
        "deck": {
            "title": "Unlabeled Versus",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "versus",
                    "title": "A or B",
                    "support_points": [{"point": "left"}, {"point": "right"}],
                }
            ],
        }
    }
    findings = validate_visuals({"deck_plan": plan})
    assert any(
        f.check == "layout" and "visual.columns" in f.detail and f.severity == "warn"
        for f in findings
    )

    html_out = tmp_path / "deck.html"
    result = render_html_deck(plan, tmp_path, html_out, language="en", evidence=None)
    doc = html_out.read_text(encoding="utf-8")
    assert any("visual.columns" in f.detail for f in result.findings)
    assert '<div class="versus-header kicker">A</div>' in doc
    assert '<div class="versus-header kicker">B</div>' in doc


def test_timeline_detail_overflow_reports_finding(tmp_path: Path):
    plan = {
        "deck": {
            "title": "Wordy Timeline",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "timeline",
                    "title": "Three Steps",
                    "support_points": [
                        {"point": "Step", "detail": "overflow " * 120},
                        {"point": "Step 2", "detail": "short"},
                        {"point": "Step 3", "detail": "short"},
                    ],
                }
            ],
        }
    }
    out = tmp_path / "deck.pptx"
    result = render_deck(plan, tmp_path, out, language="en", evidence=None)
    assert any("does not fit the card" in f.detail for f in result.findings)


def test_timeline_excess_milestones_warning(tmp_path: Path):
    plan = {
        "deck": {
            "title": "Crowded Timeline",
            "pages": [
                {
                    "id": "P01",
                    "page_role": "timeline",
                    "title": "Five Steps",
                    "support_points": [
                        {"point": "Step 1"},
                        {"point": "Step 2"},
                        {"point": "Step 3"},
                        {"point": "Step 4"},
                        {"point": "Step 5"},
                    ],
                }
            ],
        }
    }
    findings = validate_visuals({"deck_plan": plan})
    warns = [f for f in findings if f.check == "layout" and "recommended 3-4" in f.detail]
    assert len(warns) == 1
