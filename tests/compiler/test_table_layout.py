from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from docx_harness.table_layout import (
    TableCellMeasurement,
    TableLayoutPolicy,
    apply_table_layout,
    compute_table_layout,
    measure_table_cells,
)


def test_allocates_equal_and_weighted_column_widths_deterministically() -> None:
    equal = compute_table_layout(
        total_width_twips=4001,
        column_count=4,
        row_count=1,
        cells=(),
        policy=TableLayoutPolicy(column_strategy="equal"),
    )
    assert sum(equal.column_widths_twips) == 4001
    assert max(equal.column_widths_twips) - min(equal.column_widths_twips) <= 1

    weighted = compute_table_layout(
        total_width_twips=4000,
        column_count=4,
        row_count=1,
        cells=(),
        policy=TableLayoutPolicy(column_weights=(1, 2, 2, 1)),
    )
    assert sum(weighted.column_widths_twips) == 4000
    assert weighted.column_widths_twips[0] == weighted.column_widths_twips[3]
    assert weighted.column_widths_twips[1] == weighted.column_widths_twips[2]
    assert weighted.column_widths_twips[1] > weighted.column_widths_twips[0]


def test_content_strategy_gives_more_width_to_the_denser_column() -> None:
    cells = (
        TableCellMeasurement(row=0, column=0, text="术语"),
        TableCellMeasurement(row=0, column=1, text="含义"),
        TableCellMeasurement(row=1, column=0, text="DUT"),
        TableCellMeasurement(
            row=1,
            column=1,
            text="被测对象，可以是绑定物理设备的星载节点，也可以是纯仿真节点。",
        ),
    )
    content_aware = compute_table_layout(
        total_width_twips=6000,
        column_count=2,
        row_count=2,
        cells=cells,
    )
    equal = compute_table_layout(
        total_width_twips=6000,
        column_count=2,
        row_count=2,
        cells=cells,
        policy=TableLayoutPolicy(column_strategy="equal"),
    )

    assert content_aware.column_widths_twips[1] > content_aware.column_widths_twips[0]
    assert equal.column_widths_twips == (3000, 3000)


def test_full_width_merge_does_not_bias_individual_column_weights() -> None:
    plan = compute_table_layout(
        total_width_twips=6000,
        column_count=2,
        row_count=1,
        cells=(
            TableCellMeasurement(
                row=0,
                column=0,
                text="很长的跨列标题" * 20,
                column_span=2,
            ),
        ),
    )

    assert plan.column_widths_twips == (3000, 3000)


def test_horizontal_merge_uses_the_combined_width_when_estimating_height() -> None:
    text = "合并单元格自动换行" * 8
    narrow = compute_table_layout(
        total_width_twips=4000,
        column_count=2,
        row_count=1,
        cells=(TableCellMeasurement(row=0, column=0, text=text),),
    )
    merged = compute_table_layout(
        total_width_twips=4000,
        column_count=2,
        row_count=1,
        cells=(
            TableCellMeasurement(
                row=0,
                column=0,
                text=text,
                column_span=2,
            ),
        ),
    )

    assert merged.row_min_heights_twips[0] < narrow.row_min_heights_twips[0]


def test_vertical_merge_distributes_its_height_constraint_across_rows() -> None:
    plan = compute_table_layout(
        total_width_twips=2400,
        column_count=2,
        row_count=2,
        cells=(
            TableCellMeasurement(
                row=0,
                column=0,
                text="纵向合并内容" * 20,
                row_span=2,
                column_span=2,
            ),
        ),
    )

    assert sum(plan.row_min_heights_twips) > 760
    assert abs(plan.row_min_heights_twips[0] - plan.row_min_heights_twips[1]) <= 1


def test_discovers_and_applies_horizontal_and_vertical_word_merges() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "横向合并标题"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "纵向合并字段"
    table.cell(1, 1).text = "值一"
    table.cell(1, 2).text = "值二"
    table.cell(2, 1).text = "值三"
    table.cell(2, 2).text = "值四"

    cells = measure_table_cells(table)
    assert any(cell.row == 0 and cell.column_span == 2 for cell in cells)
    assert any(cell.row == 1 and cell.row_span == 2 for cell in cells)

    plan = apply_table_layout(table, 6000)
    assert sum(plan.column_widths_twips) == 6000
    assert plan.column_widths_twips[0] > plan.column_widths_twips[2]
    merged_tc = table._tbl.tr_lst[0].tc_lst[0]
    merged_width = sum(plan.column_widths_twips[:2])
    assert merged_tc.get_or_add_tcPr().get_or_add_tcW().get(qn("w:w")) == str(merged_width)
    for row, expected_height in zip(table.rows, plan.row_min_heights_twips, strict=True):
        height = row._tr.get_or_add_trPr().find(qn("w:trHeight"))
        assert height is not None
        assert height.get(qn("w:hRule")) == "atLeast"
        assert int(height.get(qn("w:val"))) == expected_height


def test_measures_native_office_math_text_for_table_layout() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=1)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph._p.append(
        parse_xml(f"<m:oMath {nsdecls('m')}><m:r><m:t>D_forward</m:t></m:r></m:oMath>")
    )

    cells = measure_table_cells(table)

    assert cells[0].text == "D_forward"
