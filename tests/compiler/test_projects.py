import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from docx_harness.components import clone_table_component
from docx_harness.errors import DocumentError
from docx_harness.project import (
    init_project,
    init_workspace,
    load_project,
    load_project_runtime,
    render_project,
    validate_project,
)
from docx_harness.style_import import extract_template

REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
PROJECT_MAINTENANCE_GUIDE = REPOSITORY_ROOT / "docs" / "compiler" / "project-maintenance.md"


def test_parent_owns_the_manual_project_maintenance_guide() -> None:
    guide = PROJECT_MAINTENANCE_GUIDE.read_text(encoding="utf-8")

    for instruction in (
        "documents/ 中的 Markdown/MyST 和手工维护的 Draw.io",
        "不要手工修改 `build/`",
        "drawio-file",
        "source",
        "target",
        "validate-project projects/<project>",
        "render-project projects/<project>",
        "提交前检查单",
    ):
        assert instruction in guide


def test_initializes_and_renders_standard_project(tmp_path: Path) -> None:
    workspace = init_workspace(tmp_path)
    project = init_project("test-reports", workspace=workspace)
    config = load_project(project)

    assert config.template.exists()
    assert config.extensions.exists()
    assert config.style_manifest.exists()
    assert (project / ".git").exists()
    assert (project / ".gitignore").exists()
    assert "/.codegraph/" in (project / ".gitignore").read_text(encoding="utf-8")
    assert (project / ".workspace" / "CONTEXT.md").exists()
    readme = (project / "README.md").read_text(encoding="utf-8")
    assert "docs/project-maintenance.md" in readme
    assert "must not be edited as sources" in readme
    output = render_project(project, "documents/example.md")
    rendered = Document(output)
    assert output == project / "build" / "example.docx"
    assert len(rendered.tables) == 1
    assert rendered.tables[0].cell(1, 1).text == "TC-EXAMPLE-001"


def test_refuses_to_overwrite_existing_project(tmp_path: Path) -> None:
    workspace = init_workspace(tmp_path)
    init_project("test-reports", workspace=workspace)

    with pytest.raises(DocumentError, match="not empty"):
        init_project("test-reports", workspace=workspace)


def test_extracts_styles_and_table_components(tmp_path: Path) -> None:
    source = tmp_path / "style-source.docx"
    source_document = Document()
    source_document.styles.add_style("Customer Body", WD_STYLE_TYPE.PARAGRAPH)
    table = source_document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    source_document.save(source)

    template = tmp_path / "base.docx"
    manifest = tmp_path / "style-manifest.json"
    components = tmp_path / "components"
    extract_template(source, template, manifest=manifest, components=components)

    extracted = Document(template)
    component = Document(components / "table-001.docx")
    inventory = json.loads(manifest.read_text(encoding="utf-8"))
    assert "Customer Body" in extracted.styles
    assert len(extracted.tables) == 0
    assert len(component.tables) == 1
    assert inventory["schema"] == "docx-harness/style-manifest/v1"
    assert inventory["tables"][0]["rows"] == 2

    output_document = Document(template)
    cloned = clone_table_component(output_document, components / "table-001.docx")
    cloned.cell(1, 1).text = "Project value"
    cloned_output = tmp_path / "cloned.docx"
    output_document.save(cloned_output)
    assert Document(cloned_output).tables[0].cell(1, 1).text == "Project value"


def test_initializes_project_from_user_docx(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Prototype"
    document.save(source)

    workspace = init_workspace(tmp_path)
    project = init_project("customer-specs", workspace=workspace, from_docx=source)

    assert (project / "references" / "style-source.docx").exists()
    assert (project / "components" / "table-001.docx").exists()
    assert Document(project / "templates" / "base.docx").tables == []


def test_managed_project_rejects_external_sources_and_outputs(tmp_path: Path) -> None:
    project = init_project("managed-reports", workspace=tmp_path)
    external = tmp_path / "external.md"
    external.write_text("External", encoding="utf-8")

    with pytest.raises(DocumentError, match="source must be inside"):
        render_project(project, external)
    with pytest.raises(DocumentError, match="output must be inside"):
        render_project(project, "documents/example.md", output=tmp_path / "outside.docx")


def test_project_relative_modules_and_lifecycle_hooks_are_isolated(tmp_path: Path) -> None:
    first = init_project("first-project", workspace=tmp_path)
    second = init_project("second-project", workspace=tmp_path)
    for project, marker in ((first, "FIRST"), (second, "SECOND")):
        (project / "project_hooks.py").write_text(
            f'MARKER = "{marker}"\n\n'
            "def finalize(document, context):\n"
            "    del context\n"
            "    document.add_paragraph(MARKER, style='DR Body')\n",
            encoding="utf-8",
        )
        (project / "extensions.py").write_text(
            "from docx_harness.extension import default_registry\n"
            "from .project_hooks import finalize\n\n"
            "def create_registry():\n"
            "    return default_registry()\n\n"
            "finalize_document = finalize\n",
            encoding="utf-8",
        )

    assert load_project_runtime(first).module.__package__ != (
        load_project_runtime(second).module.__package__
    )
    first_output = render_project(first, "documents/example.md")
    second_output = render_project(second, "documents/example.md")
    assert Document(first_output).paragraphs[-1].text == "FIRST"
    assert Document(second_output).paragraphs[-1].text == "SECOND"


def test_project_runtime_reloads_when_local_python_changes(tmp_path: Path) -> None:
    project = init_project("reload-project", workspace=tmp_path)
    helper = project / "project_hooks.py"
    helper.write_text('MARKER = "FIRST"\n', encoding="utf-8")
    (project / "extensions.py").write_text(
        "from docx_harness.extension import default_registry\n"
        "from .project_hooks import MARKER\n\n"
        "def create_registry():\n"
        "    return default_registry()\n",
        encoding="utf-8",
    )

    first = load_project_runtime(project)
    helper.write_text('MARKER = "UPDATED"\n', encoding="utf-8")
    second = load_project_runtime(project)

    assert first.module.MARKER == "FIRST"
    assert second.module.MARKER == "UPDATED"
    assert first.module.__package__ != second.module.__package__


def test_project_configuration_rejects_unknown_fields(tmp_path: Path) -> None:
    project = init_project("schema-project", workspace=tmp_path)
    config_path = project / "project.toml"
    config_path.write_text(
        config_path.read_text(encoding="utf-8") + '\nsourcee = "documents/typo.md"\n',
        encoding="utf-8",
    )

    with pytest.raises(DocumentError, match="unknown project configuration field: sourcee"):
        load_project(project)


def test_style_manifest_is_optional_evidence(tmp_path: Path) -> None:
    project = init_project("manifest-project", workspace=tmp_path)
    config_path = project / "project.toml"
    lines = [
        line
        for line in config_path.read_text(encoding="utf-8").splitlines()
        if not line.startswith("style_manifest =")
    ]
    config_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    (project / "style-manifest.json").unlink()

    config = load_project(project)
    report = validate_project(project)

    assert config.style_manifest is None
    assert "optional style manifest" not in report.checks


def test_validate_project_rejects_unknown_style_manifest_schema(tmp_path: Path) -> None:
    project = init_project("invalid-manifest-project", workspace=tmp_path)
    manifest_path = project / "style-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["schema"] = "example.invalid/style-manifest/v1"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    with pytest.raises(DocumentError, match="unsupported project style manifest schema"):
        validate_project(project)


def test_validate_project_checks_managed_git_root(tmp_path: Path) -> None:
    project = init_project("git-project", workspace=tmp_path)
    (project / ".git").rename(project / ".git-detached")

    assert load_project(project).root == project
    assert render_project(project).exists()
    with pytest.raises(DocumentError, match="not an independent Git repository"):
        validate_project(project)


def test_validate_project_checks_template_manifest_and_source(tmp_path: Path) -> None:
    project = init_project("validation-project", workspace=tmp_path)
    (project / "style-manifest.json").write_text(
        '{"schema": "unsupported/style-manifest"}\n', encoding="utf-8"
    )
    with pytest.raises(DocumentError, match="unsupported project style manifest schema"):
        validate_project(project)

    (project / "project.toml").write_text(
        (project / "project.toml")
        .read_text(encoding="utf-8")
        .replace('style_manifest = "style-manifest.json"\n', ""),
        encoding="utf-8",
    )
    (project / "templates" / "base.docx").write_text("not a DOCX", encoding="utf-8")
    with pytest.raises(DocumentError, match="cannot open project template"):
        validate_project(project)


def test_validate_project_parses_the_default_source(tmp_path: Path) -> None:
    project = init_project("source-project", workspace=tmp_path)
    (project / "documents" / "example.md").write_text(
        "# Invalid formula\n\n$$\\definitelyUnsupported{x}$$\n", encoding="utf-8"
    )

    with pytest.raises(DocumentError, match="unsupported TeX command"):
        validate_project(project)
