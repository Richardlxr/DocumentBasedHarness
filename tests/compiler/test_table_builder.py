from dataclasses import replace

import pytest
from docx import Document
from docx.oxml.ns import qn

from docx_harness.table_builder import (
    TableCellSpec,
    TableRowSpec,
    TableSpec,
    build_table,
)
from docx_harness.table_format import (
    PLAIN_COMPACT_TABLE_PROFILE,
    TableBorderProfile,
    TableCellMargins,
    TableFormatProfile,
)
from docx_harness.table_layout import TableLayoutPolicy, measure_table_cells


def _text(value: str):
    def write(cell) -> None:
        cell.text = value

    return write


def test_builder_creates_horizontal_and_vertical_merges() -> None:
    document = Document()
    spec = TableSpec(
        columns=3,
        rows=(
            TableRowSpec(
                cells=(TableCellSpec(column=0, column_span=3, write=_text("Title")),),
                repeat_header=True,
            ),
            TableRowSpec(
                cells=(
                    TableCellSpec(column=0, row_span=2, write=_text("Field")),
                    TableCellSpec(column=1, write=_text("A")),
                    TableCellSpec(column=2, write=_text("B")),
                )
            ),
            TableRowSpec(
                cells=(
                    TableCellSpec(column=1, write=_text("C")),
                    TableCellSpec(column=2, write=_text("D")),
                )
            ),
        ),
    )

    built = build_table(document, spec, available_width_twips=6000)
    cells = measure_table_cells(built.table)

    assert built.layout is not None
    assert any(cell.column_span == 3 for cell in cells)
    assert any(cell.row_span == 2 for cell in cells)
    assert built.table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None


def test_builder_rejects_overlapping_and_out_of_bounds_cells() -> None:
    document = Document()
    overlap = TableSpec(
        columns=2,
        rows=(
            TableRowSpec(
                cells=(
                    TableCellSpec(column=0, column_span=2, write=_text("A")),
                    TableCellSpec(column=1, write=_text("B")),
                )
            ),
        ),
    )
    outside = TableSpec(
        columns=2,
        rows=(TableRowSpec(cells=(TableCellSpec(column=2, write=_text("A")),)),),
    )

    with pytest.raises(ValueError, match="overlap"):
        build_table(document, overlap, available_width_twips=4000)
    with pytest.raises(ValueError, match="outside"):
        build_table(document, outside, available_width_twips=4000)


def test_profile_components_can_be_overridden_independently() -> None:
    document = Document()
    profile = TableFormatProfile(
        borders=TableBorderProfile(
            color="123456",
            outer_width_pt=2,
            inner_width_pt=1,
        ),
        cell_margins=TableCellMargins(
            top_pt=0.5,
            bottom_pt=1,
            start_pt=1.5,
            end_pt=2,
        ),
        layout=TableLayoutPolicy(column_strategy="equal"),
        fixed_layout=False,
        prevent_row_split=False,
    )
    spec = TableSpec(
        columns=2,
        rows=(
            TableRowSpec(
                cells=(
                    TableCellSpec(column=0, write=_text("A")),
                    TableCellSpec(column=1, write=_text("B")),
                )
            ),
        ),
    )

    built = build_table(document, spec, available_width_twips=4000, profile=profile)
    table = built.table
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    margins = table.cell(0, 0)._tc.get_or_add_tcPr().find(qn("w:tcMar"))

    assert table._tbl.tblPr.find(qn("w:tblLayout")) is None
    assert table.rows[0]._tr.trPr.find(qn("w:cantSplit")) is None
    assert borders.find(qn("w:top")).get(qn("w:sz")) == "16"
    assert borders.find(qn("w:insideH")).get(qn("w:sz")) == "8"
    assert borders.find(qn("w:top")).get(qn("w:color")) == "123456"
    assert margins.find(qn("w:top")).get(qn("w:w")) == "10"
    assert margins.find(qn("w:bottom")).get(qn("w:w")) == "20"
    assert margins.find(qn("w:start")).get(qn("w:w")) == "30"
    assert margins.find(qn("w:end")).get(qn("w:w")) == "40"
    assert built.layout.column_widths_twips == (2000, 2000)


def test_profile_can_opt_out_of_builder_formatting_and_layout() -> None:
    document = Document()
    profile = TableFormatProfile(
        borders=None,
        cell_margins=None,
        layout=None,
        fixed_layout=False,
        prevent_row_split=False,
    )
    spec = TableSpec(
        columns=1,
        rows=(TableRowSpec(cells=(TableCellSpec(column=0, write=_text("A")),)),),
    )

    built = build_table(document, spec, available_width_twips=2000, profile=profile)
    table = built.table

    assert built.layout is None
    assert table._tbl.tblPr.find(qn("w:tblLayout")) is None
    assert table._tbl.tblPr.find(qn("w:tblBorders")) is None
    assert table.cell(0, 0)._tc.get_or_add_tcPr().find(qn("w:tcMar")) is None
    row_properties = table.rows[0]._tr.trPr
    assert row_properties is None or row_properties.find(qn("w:cantSplit")) is None
    assert row_properties is None or row_properties.find(qn("w:trHeight")) is None


def test_default_profile_is_frozen_but_easy_to_derive() -> None:
    derived = replace(
        PLAIN_COMPACT_TABLE_PROFILE,
        layout=replace(
            PLAIN_COMPACT_TABLE_PROFILE.layout,
            column_weights=(1, 3),
        ),
    )

    assert derived.layout.column_weights == (1, 3)
    assert PLAIN_COMPACT_TABLE_PROFILE.layout.column_weights is None
