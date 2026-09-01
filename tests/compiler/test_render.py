from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from docx_harness.compiler import compile_text
from docx_harness.errors import DocumentError
from docx_harness.lifecycle import RenderHooks
from docx_harness.renderers.styles import (
    BODY_TEXT_STYLE,
    CODE_BLOCK_STYLE,
    CODE_CHARACTER_STYLE,
    MATH_BLOCK_STYLE,
    TABLE_CODE_CHARACTER_STYLE,
    TABLE_FIELD_LABEL_STYLE,
    TABLE_HEADING_STYLE,
    TABLE_MATH_BLOCK_STYLE,
    TABLE_TEXT_STYLE,
    TEST_CASE_TABLE_STYLE,
    TEST_CASE_TITLE_STYLE,
)
from docx_harness.table_format import PLAIN_COMPACT_TABLE_PROFILE, TableBorderProfile
from docx_harness.table_layout import TableLayoutPolicy
from docx_harness.template import create_template

SOURCE = """\
# Test report

Ordinary text with **bold content**.

:::{test-case} Successful login
:id: TC-AUTH-001
:priority: P0

The user enters valid `credentials`.
:::
"""


def test_generates_template_from_python(tmp_path: Path) -> None:
    template = create_template(tmp_path / "base.docx")
    document = Document(template)

    assert "DR Body" in document.styles
    assert BODY_TEXT_STYLE in document.styles
    assert TEST_CASE_TABLE_STYLE in document.styles
    assert document.styles[CODE_BLOCK_STYLE].font.size == Pt(9.5)
    assert document.styles[CODE_CHARACTER_STYLE].font.size == (
        document.styles[BODY_TEXT_STYLE].font.size
    )
    assert abs(document.sections[0].page_width - Cm(21.0)) < Cm(0.01)
    assert abs(document.sections[0].page_height - Cm(29.7)) < Cm(0.01)


def test_renders_test_case_as_fixed_layout_table(tmp_path: Path) -> None:
    template = create_template(tmp_path / "base.docx")
    output = compile_text(
        SOURCE,
        tmp_path / "report.docx",
        source="report.md",
        template=template,
    )
    document = Document(output)

    assert document.paragraphs[0].text == "Test report"
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.style.name == TEST_CASE_TABLE_STYLE
    assert table.cell(0, 0).text == "Successful login"
    assert table.cell(1, 1).text == "TC-AUTH-001"
    assert table.cell(1, 4).text == "P0"
    assert document.paragraphs[1].style.name == BODY_TEXT_STYLE
    assert table.cell(0, 0).paragraphs[0].style.name == TEST_CASE_TITLE_STYLE
    assert table.cell(0, 0).paragraphs[0].style.paragraph_format.alignment == (
        WD_ALIGN_PARAGRAPH.CENTER
    )
    assert table.cell(1, 0).paragraphs[0].style.name == TABLE_FIELD_LABEL_STYLE
    assert table.cell(1, 0).paragraphs[0].style.paragraph_format.alignment == (
        WD_ALIGN_PARAGRAPH.LEFT
    )
    assert table.cell(1, 1).paragraphs[0].style.name == TABLE_TEXT_STYLE
    assert table.cell(1, 3).paragraphs[0].style.name == TABLE_FIELD_LABEL_STYLE
    assert table.cell(2, 0).paragraphs[0].style.name == TABLE_TEXT_STYLE
    assert table.rows[0]._tr.trPr.find(qn("w:cantSplit")) is not None
    assert table.rows[1]._tr.trPr.find(qn("w:cantSplit")) is not None
    assert table.rows[2]._tr.trPr.find(qn("w:cantSplit")) is None
    code_run = next(run for run in table.cell(2, 0).paragraphs[0].runs if run.text == "credentials")
    assert code_run.style.name == TABLE_CODE_CHARACTER_STYLE
    grid_widths = [
        int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.findall(qn("w:gridCol"))
    ]
    assert len(grid_widths) == 6
    assert sum(grid_widths) > 0

    with ZipFile(output) as package:
        document_xml = package.read("word/document.xml")
    assert b'w:tblLayout w:type="fixed"' in document_xml
    assert b"TC-AUTH-001" in document_xml
    assert b"<w:shd" not in table._tbl.xml.encode()


def test_starts_document_body_with_ooxml_page_break(tmp_path: Path) -> None:
    output = compile_text(
        "# Cover\n\n2026年8月\n\n:::{document-body-start}\n:::\n\n## Content\n",
        tmp_path / "paged.docx",
        source="paged.md",
    )

    document = Document(output)
    assert [paragraph.text for paragraph in document.paragraphs] == [
        "Cover",
        "2026年8月",
        "",
        "Content",
    ]
    with ZipFile(output) as package:
        document_xml = package.read("word/document.xml")
    assert document_xml.count(b'<w:br w:type="page"/>') == 1


def test_renders_document_cover_with_centered_title_and_bottom_date(tmp_path: Path) -> None:
    template = create_template(tmp_path / "base.docx", preset="cn-official")
    output = compile_text(
        "# Cover\n\n:::{document-cover}\n:date: 编制日期：2026年8月\n:::\n\n## Content\n",
        tmp_path / "covered.docx",
        source="covered.md",
        template=template,
    )

    document = Document(output)
    title, page_break, content = document.paragraphs
    section = document.sections[0]
    title_line_height = title.paragraph_format.line_spacing
    title_center = section.top_margin + title.paragraph_format.space_before + title_line_height / 2
    date = section.first_page_footer.paragraphs[0]

    assert title.text == "Cover"
    assert date.text == "编制日期：2026年8月"
    assert page_break.text == ""
    assert content.text == "Content"
    assert title.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert date.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert abs(title_center - section.page_height / 2) < Pt(0.1)
    assert section.different_first_page_header_footer
    with ZipFile(output) as package:
        document_xml = package.read("word/document.xml")
    assert document_xml.count(b'<w:br w:type="page"/>') == 1


def test_restarts_each_ordered_list(tmp_path: Path) -> None:
    source = """\
1. First item
2. Second item

Between lists.

1. New first item
2. New second item
"""
    output = compile_text(source, tmp_path / "lists.docx", source="lists.md")
    document = Document(output)
    numbered = [
        paragraph for paragraph in document.paragraphs if paragraph.style.name == "DR Number"
    ]
    num_ids = [int(paragraph._p.pPr.numPr.numId.val) for paragraph in numbered]

    assert len(num_ids) == 4
    assert num_ids[0] == num_ids[1]
    assert num_ids[2] == num_ids[3]
    assert num_ids[0] != num_ids[2]
    numbering = document.part.numbering_part.element
    for target_num_id in set(num_ids):
        instance = next(
            num
            for num in numbering.findall(qn("w:num"))
            if int(num.get(qn("w:numId"))) == target_num_id
        )
        level_override = instance.find(qn("w:lvlOverride"))
        assert level_override is not None
        start_override = level_override.find(qn("w:startOverride"))
        assert start_override is not None
        assert start_override.get(qn("w:val")) == "1"


def test_renders_tex_as_native_editable_office_math(tmp_path: Path) -> None:
    source = r"""正文中的行内公式 $D_i=T_{2,i}-T_{0,i}$ 保持在原段落中。

$$
\bar{x}=\frac{1}{n}\sum_{i=1}^{n}x_i
$$

:::{test-case} 公式上下文
:id: TC-MATH-001

$$
s=\sqrt{\frac{\sum_{i=1}^{n}(x_i-\bar{x})^2}{n-1}}
$$

公式后的正文必须另起一段，并保留行内公式 $T_i$。
:::
"""
    output = compile_text(source, tmp_path / "math.docx", source="math.md")
    document = Document(output)

    inline_paragraph = document.paragraphs[0]
    display_paragraph = document.paragraphs[1]
    assert inline_paragraph._p.find(qn("m:oMath")) is not None
    assert display_paragraph.style.name == MATH_BLOCK_STYLE
    assert display_paragraph._p.find(qn("m:oMath")) is not None
    assert display_paragraph.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER

    content_cell = document.tables[0].cell(2, 0)
    assert content_cell.paragraphs[0].style.name == TABLE_MATH_BLOCK_STYLE
    assert content_cell.paragraphs[0]._p.find(qn("m:oMath")) is not None
    assert content_cell.paragraphs[1].text == "公式后的正文必须另起一段，并保留行内公式 。"
    assert content_cell.paragraphs[1]._p.find(qn("m:oMath")) is not None

    with ZipFile(output) as package:
        document_xml = package.read("word/document.xml")
    assert document_xml.count(b"<m:oMath>") == 4
    assert b"<m:f>" in document_xml
    assert b"<m:rad>" in document_xml


def test_renders_semantic_test_case_sections_as_merge_aware_rows(tmp_path: Path) -> None:
    source = """\
:::{test-case} Structured case
:id: TC-STRUCT-001
:priority: P1

**Purpose:** Verify a reusable structured table.

**Preconditions:** The environment is ready.

**Steps:**

1. Capture the baseline.
2. Measure inline $T_i$ and retain the raw evidence.
3. Restore the environment.
4. Verify the restored baseline.
5. Archive the evidence.

**Expected result:** Every step completes and the environment is restored.
:::
"""
    output = compile_text(source, tmp_path / "structured-case.docx", source="structured.md")
    table = Document(output).tables[0]

    assert len(table.columns) == 6
    assert len(table.rows) == 10
    assert table.cell(0, 0).text == "Structured case"
    assert table.cell(1, 0).text == "用例 ID"
    assert table.cell(1, 1).text == "TC-STRUCT-001"
    assert table.cell(1, 3).text == "优先级"
    assert table.cell(1, 4).text == "P1"

    assert table.cell(2, 0).text == "Purpose"
    assert table.cell(2, 1)._tc is table.cell(2, 5)._tc
    assert table.cell(4, 0).text == "Steps"
    assert table.cell(4, 0)._tc is table.cell(5, 0)._tc
    assert table.cell(4, 0)._tc is table.cell(6, 0)._tc
    assert table.cell(4, 0)._tc is table.cell(7, 0)._tc
    assert table.cell(4, 0)._tc is table.cell(8, 0)._tc
    assert [table.cell(row, 1).text for row in (4, 5, 6, 7, 8)] == [
        "1.",
        "2.",
        "3.",
        "4.",
        "5.",
    ]
    assert all(table.cell(row, 2)._tc is table.cell(row, 5)._tc for row in (4, 5, 6, 7, 8))
    assert table.cell(5, 2).paragraphs[0]._p.find(qn("m:oMath")) is not None
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert table.rows[1]._tr.trPr.find(qn("w:tblHeader")) is not None
    assert table.rows[2]._tr.trPr.find(qn("w:tblHeader")) is None
    assert table.cell(0, 0)._tc.tcPr.gridSpan.val == 6
    assert table.cell(1, 1)._tc.tcPr.gridSpan.val == 2
    assert table.cell(1, 4)._tc.tcPr.gridSpan.val == 2
    assert "<w:vMerge" in table._tbl.xml
    assert "<w:shd" not in table._tbl.xml


def test_prevents_normal_table_rows_from_splitting_across_pages(tmp_path: Path) -> None:
    source = """\
| Field A | Field B | Field C | Field D |
| --- | --- | --- | --- |
| A | B | C | D |
"""
    output = compile_text(source, tmp_path / "table.docx", source="table.md")
    table = Document(output).tables[0]

    assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in table.rows)
    assert all(
        paragraph.style.name in {TABLE_HEADING_STYLE, TABLE_TEXT_STYLE}
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    for edge in ("top", "start", "bottom", "end"):
        border = borders.find(qn(f"w:{edge}"))
        assert border is not None
        assert border.get(qn("w:sz")) == "12"
        assert border.get(qn("w:color")) == "000000"
    for edge in ("insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        assert border is not None
        assert border.get(qn("w:sz")) == "4"
        assert border.get(qn("w:color")) == "000000"
    assert "<w:shd" not in table._tbl.xml
    assert all(
        paragraph.style.name == TABLE_HEADING_STYLE
        for paragraph in table.rows[0].cells[0].paragraphs
    )
    assert all(
        cell.paragraphs[0].style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
        for cell in table.rows[0].cells
    )
    grid_widths = [
        int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.findall(qn("w:gridCol"))
    ]
    assert len(grid_widths) == 4
    assert max(grid_widths) - min(grid_widths) <= 1
    for row in table.rows:
        height = row._tr.trPr.find(qn("w:trHeight"))
        assert height is not None
        assert height.get(qn("w:hRule")) == "atLeast"
        assert int(height.get(qn("w:val"))) > 0
        for cell in row.cells:
            margins = cell._tc.get_or_add_tcPr().find(qn("w:tcMar"))
            assert margins is not None
            assert margins.find(qn("w:top")).get(qn("w:w")) == "40"
            assert margins.find(qn("w:bottom")).get(qn("w:w")) == "40"
            assert margins.find(qn("w:start")).get(qn("w:w")) == "60"
            assert margins.find(qn("w:end")).get(qn("w:w")) == "60"
            assert cell._tc.get_or_add_tcPr().find(qn("w:noWrap")) is None


def test_compile_accepts_a_renderer_wide_table_profile(tmp_path: Path) -> None:
    source = """\
| Short | Much longer content column |
| --- | --- |
| A | This column normally receives more width. |
"""
    profile = replace(
        PLAIN_COMPACT_TABLE_PROFILE,
        borders=TableBorderProfile(outer_width_pt=2, inner_width_pt=1),
        layout=TableLayoutPolicy(column_strategy="equal"),
    )
    output = compile_text(
        source,
        tmp_path / "custom-table-profile.docx",
        source="custom-table-profile.md",
        table_profile=profile,
    )
    table = Document(output).tables[0]
    grid_widths = [
        int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.findall(qn("w:gridCol"))
    ]
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))

    assert max(grid_widths) - min(grid_widths) <= 1
    assert borders.find(qn("w:top")).get(qn("w:sz")) == "16"
    assert borders.find(qn("w:insideH")).get(qn("w:sz")) == "8"


def test_runs_document_lifecycle_hooks_before_atomic_save(tmp_path: Path) -> None:
    events = []

    def configure(document, context) -> None:
        events.append(("configure", context.source))
        document.add_paragraph("configured", style="DR Body")

    def finalize(document, context) -> None:
        events.append(("finalize", context.output.name))
        document.add_paragraph("finalized", style="DR Body")

    def validate(document, context) -> None:
        events.append(("validate", context.project_root))
        assert document.paragraphs[-1].text == "finalized"

    output = compile_text(
        "Rendered body.",
        tmp_path / "hooked.docx",
        source="hooked.md",
        hooks=RenderHooks(configure, finalize, validate),
        project_root=tmp_path,
    )

    assert [event[0] for event in events] == ["configure", "finalize", "validate"]
    assert [paragraph.text for paragraph in Document(output).paragraphs] == [
        "configured",
        "Rendered body.",
        "finalized",
    ]


def test_failed_render_preserves_previous_output_and_diagrams(tmp_path: Path) -> None:
    destination = tmp_path / "report.docx"
    destination.write_bytes(b"previous-output")
    diagram_dir = tmp_path / "diagrams" / "report"
    diagram_dir.mkdir(parents=True)
    retained = diagram_dir / "diagram-001.drawio"
    retained.write_text("previous-diagram", encoding="utf-8")

    def fail(document, context) -> None:
        del document, context
        raise DocumentError("finalization failed")

    with pytest.raises(DocumentError, match="finalization failed"):
        compile_text(
            "New body.",
            destination,
            hooks=RenderHooks(finalize_document=fail),
        )

    assert destination.read_bytes() == b"previous-output"
    assert retained.read_text(encoding="utf-8") == "previous-diagram"


def test_relative_links_render_and_internal_links_fail_explicitly(tmp_path: Path) -> None:
    output = compile_text(
        "[Reference](docs/reference.md)",
        tmp_path / "relative-link.docx",
    )
    relationships = Document(output).part.rels.values()
    assert any(relation.target_ref == "docs/reference.md" for relation in relationships)

    with pytest.raises(DocumentError, match="internal document links"):
        compile_text("[Section](#section)", tmp_path / "internal-link.docx")
